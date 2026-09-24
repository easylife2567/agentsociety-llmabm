#!/usr/bin/env python3
"""
独立批量打标脚本 - 8 并发版本
不参考 DeepSeek 已有标签，独立判定三列
"""
import os, sys, json, time, dotenv, openpyxl
from openpyxl import Workbook
from openai import OpenAI
from concurrent.futures import ThreadPoolExecutor, as_completed

dotenv.load_dotenv("/Users/easylife/Project/AgentSociety/.env")
API_KEY = os.environ["AGENTSOCIETY_LLM_API_KEY"]
API_BASE = os.environ["AGENTSOCIETY_LLM_API_BASE"]
MODEL = os.environ["AGENTSOCIETY_LLM_MODEL"]

SRC = "/Users/easylife/Project/AgentSociety/抖音微博小红书-全量已打标.xlsx"
OUT = "/Users/easylife/Project/AgentSociety/抖音微博小红书-独立重打标_v1.xlsx"
PARTIAL = "/Users/easylife/Project/AgentSociety/output_partial.xlsx"
PROGRESS = "/Users/easylife/Project/AgentSociety/.label_progress.json"
WORKERS = 8

CLIENT = OpenAI(api_key=API_KEY, base_url=API_BASE, timeout=60)

SYSTEM_PROMPT = """你是一个中文社交媒体内容分析助手。请对每条帖子进行三方面判定。

【内容类别 - 六选一，不允许自创】
- 借势营销：以商品、服务、账号引流、带货、二维码、价格、链接、限时优惠等为目的的推广内容；或借某热点人物/事件为特定产品/服务背书引流
- 事件悼念讨论：对逝者、事故、灾害、公共事件的缅怀、哀悼、点烛、RIP、默哀类内容；或围绕此类事件的严肃讨论
- 教育观点讨论：知识科普、观点输出、教育理念讨论、人生感悟、有信息增量的分析类内容（非营销）
- 梗文化讨论：meme、谐音梗、段子、玩梗、跟风玩梗、无实义的网络流行语、搞笑吐槽类内容
- 其他讨论：以上四类都不符合，但仍是正常有意义的帖子（如日常分享、生活记录、提问求助、情感抒发）
- 爬取噪音：乱码、纯表情、@好友无意义内容、无意义短文本、广告刷屏残留、空内容、格式损坏等无实质信息的条目

【数据有效性 - 二选一】
- 有效：包含可用于后续研究分析的实质内容
- 无效：无实质内容、纯噪音、纯营销引流无研究价值、或内容过于残缺无法分析

【判定原则】
1. 先判断是否为噪音：内容残缺/乱码/纯@人/纯表情 → 爬取噪音，无效
2. 再判断是否为营销：带货/引流/推广 → 借势营销，无效
3. 再判断是否为悼念：对逝者/事故/灾害的缅怀 → 事件悼念讨论，有效
4. 再判断是否为玩梗：meme/段子/谐音梗/跟风玩梗 → 梗文化讨论，有效
5. 再判断是否为教育观点：知识输出/观点分析/教育理念/人生感悟 → 教育观点讨论，有效
6. 以上都不符合 → 其他讨论，根据是否有实质内容判断有效性

【输出格式 - 严格 JSON，不要任何其他文字】
{"category": "类别名", "reason": "一句话依据", "valid": "有效|无效"}"""


def label_one(row_num, title, content, media):
    user_msg = f"【平台】{media}\n【标题】{title or ''}\n【正文】{(content or '')[:2000]}"
    for attempt in range(3):
        try:
            r = CLIENT.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_msg},
                ],
                temperature=0.1, max_tokens=200,
            )
            txt = r.choices[0].message.content.strip()
            if txt.startswith("```"):
                txt = txt.split("```")[1] if "```" in txt else txt
                if txt.startswith("json"):
                    txt = txt[4:]
                txt = txt.strip()
            lo, hi = txt.find("{"), txt.rfind("}")
            if lo >= 0 and hi > lo:
                obj = json.loads(txt[lo:hi+1])
                cat = obj.get("category", "其他讨论")
                reason = obj.get("reason", "")
                valid = obj.get("valid", "有效")
                if cat not in {"借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音"}:
                    cat = "其他讨论"
                if valid not in {"有效", "无效"}:
                    valid = "有效"
                return row_num, cat, reason, valid
        except Exception as e:
            if attempt == 2:
                return row_num, "其他讨论", f"[API失败:{str(e)[:50]}]", "无效"
            time.sleep(2 ** attempt * 0.5)
    return row_num, "其他讨论", "[重试失败]", "无效"


def main():
    # 断点续跑
    start_row = 2
    done_rows = set()
    if os.path.exists(PROGRESS):
        with open(PROGRESS) as f:
            progress = json.load(f)
        start_row = progress["next_row"]
        print(f"断点续跑：从第 {start_row} 行继续")

    # 读全量数据到内存（5.2万行 × 57列，约几百MB）
    print("读取源数据...")
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows_iter = ws.iter_rows(values_only=True)
    header = list(next(rows_iter))

    idx = {name: header.index(name) if name in header else -1
           for name in ["title", "content", "full_content", "media_type", "url", "id"]}

    all_rows = []
    row_num = 1
    for row in rows_iter:
        row_num += 1
        if row is None or all(v is None for v in row):
            continue
        all_rows.append((row_num, row))
    print(f"共 {len(all_rows)} 条待处理")

    # 输出工作簿
    if os.path.exists(PARTIAL):
        out_wb = openpyxl.load_workbook(PARTIAL)
        out_ws = out_wb.active
    else:
        out_wb = Workbook()
        out_ws = out_wb.active
        out_ws.append(header + ["【新】类别", "【新】判定依据", "【新】有效性"])

    # 过滤已完成的行
    pending = [(rn, row) for rn, row in all_rows if rn >= start_row]
    print(f"待处理：{len(pending)} 条")

    t0 = time.time()
    processed = 0
    save_interval = 200
    batch_size = 80  # 每批提交80条，8并发

    for batch_start in range(0, len(pending), batch_size):
        batch = pending[batch_start:batch_start + batch_size]
        tasks = []
        with ThreadPoolExecutor(max_workers=WORKERS) as ex:
            for row_num, row in batch:
                title = row[idx["title"]] or ""
                content = row[idx["content"]] or row[idx["full_content"]] or ""
                media = row[idx["media_type"]] or ""
                tasks.append(ex.submit(label_one, row_num, title, content, media))

            results = []
            for fut in as_completed(tasks):
                results.append(fut.result())

        # 按行号排序写入
        results.sort(key=lambda x: x[0])
        for row_num, cat, reason, valid in results:
            # 找到原始行
            orig_row = None
            for rn2, r2 in batch:
                if rn2 == row_num:
                    orig_row = r2
                    break
            out_ws.append(list(orig_row) + [cat, reason, valid])
            processed += 1

        # 落盘
        if processed % save_interval == 0 or batch_start + batch_size >= len(pending):
            out_wb.save(PARTIAL)
            elapsed = time.time() - t0
            rate = processed / elapsed if elapsed > 0 else 0
            remaining = (len(pending) - processed) / rate if rate > 0 else 0
            print(f"进度 {processed}/{len(pending)} ({processed*100//len(pending)}%) | "
                  f"速度 {rate:.1f} 条/秒 | 已跑 {elapsed/60:.1f} 分钟 | "
                  f"预计剩余 {remaining/60:.1f} 分钟", flush=True)
            with open(PROGRESS, "w") as f:
                json.dump({"next_row": start_row + processed, "done": start_row + processed - 1}, f)

    # 完成
    out_wb.save(PARTIAL)
    if os.path.exists(OUT):
        os.remove(OUT)
    os.rename(PARTIAL, OUT)
    elapsed = time.time() - t0
    print(f"\n✅ 全部完成！总耗时 {elapsed/60:.1f} 分钟，输出：{OUT}")


if __name__ == "__main__":
    main()