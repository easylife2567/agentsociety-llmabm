#!/usr/bin/env python3
"""
新一轮打标（doubao / GLM × 第1轮口径）与**基线**的对齐预检（只读，不改任何数据）。

基线 = 第 1 轮原标（`baseline/抖音微博小红书-全量已打标.xlsx`）。新一轮要能与它可比，
必须同时在四个层面复刻第 1 轮，任何一层错位都会让后续的 κ 无法解释：

  A. 提示词对齐 —— batch_label_r1prompt.py 的 SYSTEM_PROMPT 与第 1 轮恢复出的原文，
     除两处 I/O 段（读批次文件→API 批量；Write 工具→JSON 数组）外必须逐字相同。
     断言方式：取【研究背景】到【输出格式/操作】之间的语义主体，要求 byte 级相等。
  B. 输入字段对齐 —— 用 build_record() 重建全部输入，校验第 1 轮记录在案的两个签名：
     补 ctx_full_content 的行数 = 1312、title 为字符串 "nan" 的占比 = 42.4%。
  C. 行号对齐 —— row_id 必须 0..52735 连续；第 1 轮的 result_batch_*.jsonl 给出其
     实际输入顺序（batch_k 覆盖 [150(k-1), 150k)），据此校验批次切分与 row_id→Excel 行号。
  D. 基线自身 —— 基线 xlsx 的 数据有效性/内容类别/判定依据 应与第 1 轮 result 文件逐行相同，
     确认「基线 = 第 1 轮产物」且未被改动，否则比对基准本身就是错的。

第 1 轮的**输入**批次（/tmp/zxf_full/batch_*.jsonl）已随 /tmp 清理，故 B 只能做签名校验
而非逐字节 diff；C/D 用第 1 轮的**输出**批次补上这层证据。

用法: python check_baseline_alignment.py [--skip-rebuild]
"""
import ast, difflib, glob, hashlib, json, os, re, sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA = f"{ROOT}/data"
NEW_SCRIPT = f"{DATA}/scripts/batch_label_r1prompt.py"
BASELINE = f"{DATA}/baseline/抖音微博小红书-全量已打标.xlsx"

# 第 1 轮提示词的恢复来源（本机 Workflow 历史脚本）与批次结果
R1_WF = os.path.expanduser(
    "~/.claude/projects/-Users-easylife-Project-----/"
    "48165eea-81a5-449c-ab05-3b11f103a1af/workflows/scripts/"
    "zxf-full-labeling-wf_4953f96e-780.js")
R1_OUT_DIR = os.path.expanduser("~/Project/数据清洗/全量标注结果")

EXPECT_ROWS = 52736
EXPECT_CTX = 1312
EXPECT_NAN_PCT = 42.4
EXPECT_BATCH_ROWS = 150

fails = []


def check(cond, label, detail=""):
    print(f"  {'✅' if cond else '❌'} {label}" + (f"  {detail}" if detail else ""))
    if not cond:
        fails.append(label)
    return cond


def section(t):
    print(f"\n{'='*78}\n{t}\n{'='*78}")


def extract_r1_prompt():
    src = open(R1_WF, encoding="utf-8").read()
    m = re.search(r"const T = `(.*?)`\n", src, re.S)
    return m.group(1).replace("\\`", "`")


def extract_new_prompt():
    tree = ast.parse(open(NEW_SCRIPT, encoding="utf-8").read())
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(
                getattr(t, "id", None) == "SYSTEM_PROMPT" for t in node.targets):
            return node.value.value
    raise SystemExit("未能在 batch_label_r1prompt.py 中定位 SYSTEM_PROMPT")


def body_of(text):
    """取【研究背景】到【输出格式】/【操作】之间的语义主体 —— 口径的核心，必须逐字相同。"""
    i = text.index("【研究背景】")
    # 第1轮原文用【操作】，新脚本改为【输出格式 - 严格 JSON 数组…】，故只匹配前缀
    j = min(x for x in (text.find("【输出格式"), text.find("【操作】")) if x > 0)
    return text[i:j]


def section_a():
    section("A. 提示词对齐（第 1 轮原文 vs batch_label_r1prompt.py）")
    if not os.path.exists(R1_WF):
        print(f"  ❌ 第 1 轮工作流脚本不在（本机历史已清理）：{R1_WF}")
        fails.append("第1轮提示词源缺失")
        return None
    old, new = extract_r1_prompt(), extract_new_prompt()
    print(f"  第1轮原文  长度={len(old):>5}  sha256={hashlib.sha256(old.encode()).hexdigest()[:16]}")
    print(f"  新脚本     长度={len(new):>5}  sha256={hashlib.sha256(new.encode()).hexdigest()[:16]}")

    ob, nb = body_of(old), body_of(new)
    check(ob == nb, "语义主体 byte 级相等（【研究背景】→【类别体系】末）",
          f"{len(ob)} 字符")

    # 把差异逐行列出，供人工确认「只有两处 I/O 段」
    diff = [l for l in difflib.unified_diff(old.splitlines(), new.splitlines(),
                                            "第1轮原文", "新脚本", lineterm="", n=0)
            if l[:1] in "+-" and l[:3] not in ("+++", "---")]
    hunks = sum(1 for l in difflib.unified_diff(old.splitlines(), new.splitlines(), n=0)
                if l.startswith("@@"))
    print(f"\n  差异块 {hunks} 处、差异行 {len(diff)} 行（应为 2 处 I/O 段）：")
    for l in diff:
        print(f"    {l[:150]}")
    check(hunks == 2, "差异恰好 2 处（I/O 段）", f"实际 {hunks}")
    return old, new


def section_bc():
    section("B/C. 输入字段与行号对齐（重建全部输入）")
    sys.path.insert(0, os.path.dirname(NEW_SCRIPT))
    # batch_label_r1prompt 在模块级读 sys.argv[1..2]，导入前先垫上占位参数
    _argv = sys.argv
    sys.argv = ["batch_label_r1prompt.py", "__probe__", "__probe__"]
    import batch_label_r1prompt as B
    sys.argv = _argv

    wb = openpyxl.load_workbook(BASELINE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    for name in ["title", "content", "full_content", "media_type", "author_name",
                 "match_sentence", "数据有效性", "内容类别", "判定依据"]:
        B.IDX[name] = header.index(name)

    recs, labels, rn = {}, [], 1
    for row in it:
        rn += 1
        if row is None or all(v is None for v in row):
            continue
        recs[rn] = B.build_record(rn - 2, row)
        labels.append((str(row[B.IDX["数据有效性"]]), str(row[B.IDX["内容类别"]]),
                       str(row[B.IDX["判定依据"]] or "")))
    wb.close()

    n = len(recs)
    check(n == EXPECT_ROWS, f"源表数据行数 = {EXPECT_ROWS}", f"实际 {n}")
    check(sorted(recs) == list(range(2, n + 2)), "Excel 行号连续无空洞")

    rids = [r["row_id"] for r in recs.values()]
    check(rids == list(range(n)), "row_id = 0..n-1 连续且与 Excel 行号差 2")

    fields = {tuple(sorted(r)) for r in recs.values()}
    check(fields == {("author", "match_sentence", "platform", "row_id", "title"),
                     ("author", "ctx_full_content", "match_sentence", "platform",
                      "row_id", "title")},
          "字段名只有两种形态（有/无 ctx_full_content）", f"{len(fields)} 种")

    nctx = sum(1 for r in recs.values() if "ctx_full_content" in r)
    check(nctx == EXPECT_CTX, f"补 ctx_full_content 的行数 = {EXPECT_CTX}",
          f"实际 {nctx}（{nctx/n*100:.2f}%）")

    nnan = sum(1 for r in recs.values() if r["title"] == "nan")
    pct = nnan / n * 100
    check(abs(pct - EXPECT_NAN_PCT) < 0.5, f'title 为字符串 "nan" 占比 ≈ {EXPECT_NAN_PCT}%',
          f"实际 {pct:.2f}%（{nnan} 行）")
    return recs, labels


def section_cd(labels):
    section("C/D. 第 1 轮输出批次：行序、批次切分、与基线逐行一致")
    files = sorted(glob.glob(os.path.join(R1_OUT_DIR, "result_batch_*.jsonl")),
                   key=lambda p: int(re.search(r"batch_(\d+)", p).group(1)))
    if not files:
        print(f"  ❌ 第 1 轮批次结果目录不存在：{R1_OUT_DIR}")
        fails.append("第1轮输出批次缺失")
        return
    r1 = []
    for f in files:
        for line in open(f, encoding="utf-8"):
            line = line.strip()
            if line:
                r1.append(json.loads(line))

    check(len(r1) == EXPECT_ROWS, f"第 1 轮批次结果共 {EXPECT_ROWS} 行", f"实际 {len(r1)}")
    check([r["row_id"] for r in r1] == list(range(len(r1))),
          "批次按 row_id 0..n-1 顺序排列（batch_k 覆盖 [150(k-1), 150k)）")
    check(len(files) == -(-EXPECT_ROWS // EXPECT_BATCH_ROWS),
          f"批次数 = ⌈{EXPECT_ROWS}/{EXPECT_BATCH_ROWS}⌉ = "
          f"{-(-EXPECT_ROWS // EXPECT_BATCH_ROWS)}", f"实际 {len(files)}")

    # D. 基线 xlsx 是否就是第 1 轮的产物（逐行）
    diff = [i for i, (r, lab) in enumerate(zip(r1, labels))
            if (r["valid"], r["category"], r["reason"]) != lab]
    check(not diff, "基线 xlsx 的 数据有效性/内容类别/判定依据 与第 1 轮批次逐行相同",
          f"{len(diff)} 行不同" + (f"，首例 行{diff[0]+2}" if diff else ""))

    L = sorted(len(r["reason"]) for r in r1)
    m = len(L)
    print(f"\n  第 1 轮输出端签名：reason 中位 {L[m//2]}  最大 {L[-1]}  "
          f"≤30字 {sum(1 for x in L if x <= 30)/m*100:.2f}%")
    print("  → 新一轮（第1轮口径）应复现同一签名；若中位/最大明显不同，说明提示词没对齐。")


def main():
    print("基线对齐预检（只读）")
    print(f"  新脚本 : {NEW_SCRIPT}")
    print(f"  基线   : {BASELINE}")
    section_a()
    if "--skip-rebuild" not in sys.argv:
        recs, labels = section_bc()
        section_cd(labels)
    print(f"\n{'='*78}")
    if fails:
        print(f"❌ 未通过 {len(fails)} 项：" + "；".join(fails))
        sys.exit(1)
    print("✅ 全部通过 —— 新一轮与基线在同一口径、同一输入字段、同一行序上可比。")


if __name__ == "__main__":
    main()
