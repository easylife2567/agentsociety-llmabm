#!/usr/bin/env python3
"""
三方标签一致性对比：DeepSeek 原标 vs doubao-seed 轮 vs GLM 轮
按行对齐读取三份 xlsx 的前三列（数据有效性/内容类别/判定依据），输出：
  - 两两总体一致率 + Cohen's kappa（类别、有效性）
  - 混淆矩阵（含 DeepSeek 特有的「存疑」值）
  - 三方同时一致率
  - 分歧样本清单（md 报告 + xlsx 明细）
用法: python compare_labels.py [--glm <GLM轮xlsx路径>]
"""
import os, sys, json, argparse, datetime
from collections import Counter
from itertools import combinations
import openpyxl
from openpyxl import Workbook

ROOT = "/Users/easylife/Project/AgentSociety"
DATA = f"{ROOT}/data"   # 打标工作总目录（2026-09-27 起，见 data/README.md）
CAT_ORDER = ["借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音", "存疑"]
VAL_ORDER = ["有效", "无效", "存疑"]


def default_paths():
    return {
        "DeepSeek": f"{DATA}/baseline/抖音微博小红书-全量已打标.xlsx",
        "doubao-seed": f"{DATA}/archive/round2_doubao-seed-2.1-lite/抖音微博小红书-独立重打标_Seed2_1_lite.xlsx",
        "GLM": f"{DATA}/archive/round3_glm-5.3-flash/抖音微博小红书-独立重打标_GLM5_3flash.xlsx",
    }


def load_labels(path):
    """返回 [(valid, cat, reason, rid, snippet)]，按行序。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    id_i, t_i, c_i = header.index("id"), header.index("title"), header.index("content")
    out = []
    for row in it:
        valid, cat, reason = str(row[0]), str(row[1]), str(row[2])
        rid = row[id_i]
        title = row[t_i] or ""
        content = row[c_i] or ""
        snippet = str(title)[:40] if str(title).strip() else str(content)[:40]
        out.append((valid, cat, reason, rid, snippet))
    wb.close()
    return out


def kappa(a, b):
    """Cohen's kappa，输入两个等长标签列表。"""
    n = len(a)
    if n == 0:
        return float("nan")
    joint = Counter(zip(a, b))
    p_o = sum(v for (x, y), v in joint.items() if x == y) / n
    ca, cb = Counter(a), Counter(b)
    p_e = sum(ca[k] * cb.get(k, 0) for k in ca) / (n * n)
    if p_e == 1:
        return 1.0
    return (p_o - p_e) / (1 - p_e)


def confusion(a, b, order):
    """返回 {x: {y: count}}，x=A标签 y=B标签。"""
    m = {x: {y: 0 for y in order} for x in order}
    for x, y in zip(a, b):
        m.setdefault(x, {y: 0 for y in order})
        m[x].setdefault(y, 0)
        m[x][y] += 1
    return m


def fmt_matrix(m, order):
    short = {c: c[:2] for c in order}
    head = "A\\B | " + " | ".join(f"{short[c]:>4}" for c in order)
    lines = [head, "-" * len(head)]
    for x in order:
        lines.append(f"{short[x]:>4} | " + " | ".join(f"{m.get(x, {}).get(y, 0):>4}" for y in order))
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--glm", default=None, help="GLM 轮 xlsx 路径（默认取 GLM5_3flash）")
    args = ap.parse_args()

    paths = default_paths()
    if args.glm:
        paths["GLM"] = args.glm

    data = {}
    for name, p in paths.items():
        if not os.path.exists(p):
            print(f"跳过缺失评审方: {name} ({p})")
            continue
        data[name] = load_labels(p)
        print(f"载入 {name}: {len(data[name])} 行")

    names = list(data)
    if len(names) < 2:
        print("至少需要两份标签文件")
        sys.exit(1)
    n = len(data[names[0]])
    for nm in names:
        assert len(data[nm]) == n, f"{nm} 行数不一致: {len(data[nm])} != {n}"
    # id 对齐抽查
    for i in [0, n // 2, n - 1]:
        ids = {str(data[nm][i][3]) for nm in names}
        assert len(ids) == 1, f"行 {i+2} id 不对齐: {ids}"
    print(f"行数与 id 对齐校验通过（n={n}）")

    # 预提取标签序列
    cats = {nm: [r[1] for r in data[nm]] for nm in names}
    vals = {nm: [r[0] for r in data[nm]] for nm in names}

    today = datetime.date.today().strftime("%Y%m%d")
    md = [f"# 三方标签一致性对比报告（{today}）\n"]
    md.append(f"- 评审方：{'、'.join(names)}")
    md.append(f"- 样本量：{n} 条")
    for nm in names:
        cc, vv = Counter(cats[nm]), Counter(vals[nm])
        md.append(f"- **{nm}** 类别分布：{'，'.join(f'{k} {v} ({v*100/n:.1f}%)' for k, v in cc.most_common())}")
        md.append(f"- **{nm}** 有效性分布：{'，'.join(f'{k} {v}' for k, v in vv.most_common())}")
    md.append("")

    xwb = Workbook()
    xwb.remove(xwb.active)

    # 两两对比
    for a, b in combinations(names, 2):
        agree_c = sum(x == y for x, y in zip(cats[a], cats[b]))
        agree_v = sum(x == y for x, y in zip(vals[a], vals[b]))
        k_c = kappa(cats[a], cats[b])
        k_v = kappa(vals[a], vals[b])
        order_c = [c for c in CAT_ORDER if c in set(cats[a]) | set(cats[b])]
        order_v = [v for v in VAL_ORDER if v in set(vals[a]) | set(vals[b])]
        md.append(f"## {a} vs {b}\n")
        md.append(f"- 类别一致率：{agree_c}/{n} = {agree_c*100/n:.2f}%，Cohen's κ = {k_c:.4f}")
        md.append(f"- 有效性一致率：{agree_v}/{n} = {agree_v*100/n:.2f}%，Cohen's κ = {k_v:.4f}")
        md.append(f"\n类别混淆矩阵（行={a}，列={b}）：\n```\n{fmt_matrix(confusion(cats[a], cats[b], order_c), order_c)}\n```")
        md.append(f"\n有效性混淆矩阵：\n```\n{fmt_matrix(confusion(vals[a], vals[b], order_v), order_v)}\n```")

        # 分歧样本（每对最多 20 条）
        dis = [(i, data[a][i], data[b][i]) for i in range(n) if cats[a][i] != cats[b][i]]
        md.append(f"- 类别分歧行数：{len(dis)}")
        md.append("\n分歧示例（行号/平台内容摘要/两方判定）：\n")
        for i, ra, rb in dis[:20]:
            md.append(f"- 行{i+2}｜{ra[4]}｜{a}: {ra[1]}（{ra[0]}）←→ {b}: {rb[1]}（{rb[0]}）")
        md.append("")

        ws = xwb.create_sheet(f"{a}-vs-{b}-混淆")
        ws.append(["A\\B"] + order_c)
        for x in order_c:
            ws.append([x] + [confusion(cats[a], cats[b], order_c).get(x, {}).get(y, 0) for y in order_c])
        wsd = xwb.create_sheet(f"{a}-vs-{b}-分歧")
        wsd.append(["行号", "id", "内容摘要", f"{a}类别", f"{a}有效性", f"{b}类别", f"{b}有效性", f"{a}依据", f"{b}依据"])
        for i, ra, rb in dis:
            wsd.append([i + 2, str(ra[3]), ra[4], ra[1], ra[0], rb[1], rb[0], ra[2], rb[2]])

    # 三方同时一致率（取存在的评审方）
    if len(names) >= 3:
        agree3 = sum(len({cats[nm][i] for nm in names}) == 1 for i in range(n))
        agree3v = sum(len({vals[nm][i] for nm in names}) == 1 for i in range(n))
        md.append("## 三方同时一致\n")
        md.append(f"- 类别三方一致：{agree3}/{n} = {agree3*100/n:.2f}%")
        md.append(f"- 有效性三方一致：{agree3v}/{n} = {agree3v*100/n:.2f}%\n")
        ws3 = xwb.create_sheet("三方分歧")
        ws3.append(["行号", "id", "内容摘要"] + [f"{nm}类别/{nm}有效性" for nm in names])
        for i in range(n):
            if len({cats[nm][i] for nm in names}) > 1:
                row = [i + 2, str(data[names[0]][i][3]), data[names[0]][i][4]]
                for nm in names:
                    row += [cats[nm][i], vals[nm][i]]
                ws3.append(row)

    # DeepSeek 存疑行单独列示
    if "DeepSeek" in data:
        dq = [i for i in range(n) if cats["DeepSeek"][i] == "存疑"]
        md.append(f"## DeepSeek 原标「存疑」行（{len(dq)} 条）\n")
        for i in dq:
            others = "，".join(f"{nm}: {cats[nm][i]}（{vals[nm][i]}）" for nm in names if nm != "DeepSeek")
            md.append(f"- 行{i+2}｜{data['DeepSeek'][i][4]}｜{others}")
        wsq = xwb.create_sheet("DeepSeek存疑行")
        wsq.append(["行号", "id", "内容摘要", "DeepSeek类别"] + [f"{nm}类别" for nm in names if nm != "DeepSeek"])
        for i in dq:
            wsq.append([i + 2, str(data["DeepSeek"][i][3]), data["DeepSeek"][i][4], "存疑"] +
                       [cats[nm][i] for nm in names if nm != "DeepSeek"])

    md_path = f"{DATA}/reports/标签一致性对比报告_{today}.md"
    xlsx_path = f"{DATA}/reports/标签一致性对比明细_{today}.xlsx"
    with open(md_path, "w") as f:
        f.write("\n".join(md))
    xwb.save(xlsx_path)
    print(f"\n✅ 报告：{md_path}\n✅ 明细：{xlsx_path}")


if __name__ == "__main__":
    main()
