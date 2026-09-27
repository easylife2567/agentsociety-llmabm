#!/usr/bin/env python3
"""把基线（原标 v4-flash）与某一轮重打标的判定并排成一张对照表，供人工复核。

表结构（前 6 列是两轮的判定，其后接源表原始数据列，共 60 列）：

  A–C  数据有效性/内容类别/判定依据 —— 原标 v4-flash（第1轮口径）
  D–F  数据有效性/内容类别/判定依据 —— 本轮模型（第1轮口径）
  G起  源表原始数据列（id、title、content…）

高亮：
  · 两轮「内容类别」不一致的行 → 整行黄底（FFFF99），即待人工复核的分歧行
  · 服务端内容过滤导致「未判定」的行 → 整行灰底（D9D9D9）
    （这 3 行不是模型分歧，单列一色以免与上者混淆，详见 data/README.md）

用 write_only 流式写出：3 百多万单元格，一次性载入内存会爆。

用法: python merge_judgments.py <slug> <模型标签>
产出: data/runs/抖音微博小红书-判定对照_v4flash_vs_<模型标签>.xlsx
"""
import os, sys
import openpyxl
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import PatternFill

ROOT = "/Users/easylife/Project/AgentSociety"
DATA = f"{ROOT}/data"
BASE = f"{DATA}/baseline/抖音微博小红书-全量已打标.xlsx"

if len(sys.argv) < 3:
    sys.exit("用法: python merge_judgments.py <slug> <模型标签>\n"
             "  例: python merge_judgments.py R1prompt_Seed2_1_lite doubao")
SLUG, LABEL = sys.argv[1], sys.argv[2]
NEW = f"{DATA}/runs/抖音微博小红书-独立重打标_{SLUG}.xlsx"
OUT = f"{DATA}/runs/抖音微博小红书-判定对照_v4flash_vs_{LABEL}.xlsx"
for p in (BASE, NEW):
    if not os.path.exists(p):
        sys.exit(f"缺文件: {p}")

FILL_DIFF = PatternFill("solid", start_color="FFFF99", end_color="FFFF99")
FILL_NA = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")

wb_b = openpyxl.load_workbook(BASE, read_only=True, data_only=True)
ws_b = wb_b[wb_b.sheetnames[0]]
it_b = ws_b.iter_rows(values_only=True)
h_b = list(next(it_b))

wb_n = openpyxl.load_workbook(NEW, read_only=True, data_only=True)
ws_n = wb_n[wb_n.sheetnames[0]]
it_n = ws_n.iter_rows(values_only=True)
h_n = list(next(it_n))

assert h_b[3:] == h_n[3:], "两份表的原始数据列不一致"
data_hdr = h_b[3:]

out = Workbook(write_only=True)
ws = out.create_sheet("判定对照")
ws.append(["数据有效性_原标v4flash", "内容类别_原标v4flash", "判定依据_原标v4flash",
           f"数据有效性_{LABEL}", f"内容类别_{LABEL}", f"判定依据_{LABEL}"] + data_hdr)

def _c(sheet, value, fill):
    c = WriteOnlyCell(sheet, value=value)
    c.fill = fill
    return c


n = n_diff = n_na = n_val_only = 0
for rb, rr in zip(it_b, it_n):
    if rb is None or rb[0] is None:
        continue
    n += 1
    assert str(rb[3]) == str(rr[3]), f"第 {n} 行 id 不对齐: {rb[3]} vs {rr[3]}"
    six = [rb[0], rb[1], rb[2], rr[0], rr[1], rr[2]]

    if str(rb[0]) == "未判定" or str(rr[0]) == "未判定":
        fill, n_na = FILL_NA, n_na + 1
    elif str(rb[1]) != str(rr[1]):
        fill, n_diff = FILL_DIFF, n_diff + 1
    else:
        fill = None
        if str(rb[0]) != str(rr[0]):
            n_val_only += 1

    if fill is None:
        ws.append(six + list(rb[3:]))
    else:
        ws.append([_c(ws, v, fill) for v in six + list(rb[3:])])


wb_b.close()
wb_n.close()
out.save(OUT)

size = os.path.getsize(OUT) / 1024 / 1024
print(f"✅ {OUT}")
print(f"   {n} 行 × {6 + len(data_hdr)} 列（6 判定列 + {len(data_hdr)} 原始数据列） | {size:.1f} MB")
print(f"   黄底（类别不一致，待人工复核）: {n_diff} 行")
print(f"   灰底（未判定，服务端内容过滤）: {n_na} 行")
print(f"   类别一致但有效性不一致（未高亮）: {n_val_only} 行")
