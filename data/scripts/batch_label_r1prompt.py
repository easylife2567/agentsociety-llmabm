#!/usr/bin/env python3
"""
以**第 1 轮（原标）提示词**重跑打标 —— 用于把「口径差异」与「判定质量差异」分离。

背景：第 1 轮提示词已于 2026-09-26 从本机 Workflow 历史脚本恢复（见 label_prompt_round1.md）。
本脚本用该提示词 + 该轮输入字段，在指定模型上重跑一遍，补上 2×2 对照中缺失的那一格：

                 第1轮口径                      后三轮口径
  原标模型      已有（=原标, agentic 150行/批）   —
  <MODEL>      **本脚本**                       已有（=第N轮, API 8条/批）

  · 与本脚本 vs 既有「后三轮口径」结果 → 同模型、同调用结构，**纯口径效应**（提示词+输入字段）
  · 本脚本 vs 原标 → 同提示词、同输入字段，**模型+调用结构效应**

提示词的改动：**仅两处 I/O 段**，其余（研究背景/判定核心标准/类别体系/存疑规则）逐字保留。
  (1) 开头「用 Read 工具读取 /tmp/zxf_full/batch_N.jsonl（…）」→「下面给你一批帖子（…）」
  (2) 结尾【操作】「用 Write 工具写入 /tmp/zxf_full/result_batch_N.jsonl … 通过 StructuredOutput
      返回计数」→【输出格式】严格 JSON 数组
差异全文见 label_prompt_round1.md §C。

输入字段**逐字复刻**第 1 轮的构建代码（含一个真实输入痕迹：源表空值经 pandas str() 变成
字符串 "nan"，第 1 轮 42.4% 的行 title 就是 "nan"；match_sentence 例外，用 fillna("") 得空串）：
  row_id(0基位置) / platform(media_type) / author(author_name[:30]) / title(title[:80])
  / match_sentence(去 <em> 标签后 [:400])；当 match_sentence <15 字时补
  ctx_full_content = full_content[:250]，为空则 content[:250]

用法: python batch_label_r1prompt.py <model_id> <slug> [workers] [batch_size]
产出: 抖音微博小红书-独立重打标_<slug>.xlsx（结构=原表57列，前三列替换为新判定）
断点: .label_results_<slug>.jsonl（schema 与 v2/v3 一致，row_num 为 Excel 1基行号）
"""
import os, re, sys, json, time, dotenv, openpyxl
from openpyxl import Workbook
from openai import OpenAI
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import Counter

dotenv.load_dotenv("/Users/easylife/Project/AgentSociety/.env")
API_KEY = os.environ["AGENTSOCIETY_LLM_API_KEY"]
API_BASE = os.environ["AGENTSOCIETY_LLM_API_BASE"]

MODEL = sys.argv[1]
SLUG = sys.argv[2]
WORKERS = int(sys.argv[3]) if len(sys.argv) > 3 else 24
BSIZE = int(sys.argv[4]) if len(sys.argv) > 4 else 8
MAX_TOKENS = int(os.environ.get("LABEL_MAX_TOKENS", "8000"))

ROOT = "/Users/easylife/Project/AgentSociety"
DATA = f"{ROOT}/data"   # 打标工作总目录（2026-09-27 起，见 data/README.md）
SRC = f"{DATA}/baseline/抖音微博小红书-全量已打标.xlsx"
OUT = f"{DATA}/runs/抖音微博小红书-独立重打标_{SLUG}.xlsx"
JSONL = f"{DATA}/runs/.label_results_{SLUG}.jsonl"
PROGRESS = f"{DATA}/runs/.label_progress_{SLUG}.json"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

print(f"[配置] 模型={MODEL} | slug={SLUG} | 并发={WORKERS} | 每次调用判定 {BSIZE} 条 | max_tokens={MAX_TOKENS}", flush=True)
print(f"[配置] 提示词=第1轮(原标)提示词，仅 I/O 段改为 API | 输入字段=第1轮原样", flush=True)
print(f"[配置] 输出={OUT}", flush=True)

# 超时可配：思考型模型单次 8 条判定实测 110–430s（doubao 思考开 176s/14k tokens 属常态），
# 180 秒会误杀本可成功的调用并触发最多 3 次重试，既空耗墙钟也重复计费 token。
# 思考型模型（doubao / GLM）用 LABEL_TIMEOUT=600。
TIMEOUT = float(os.environ.get("LABEL_TIMEOUT", "180"))
CLIENT = OpenAI(api_key=API_KEY, base_url=API_BASE, timeout=TIMEOUT)
print(f"[配置] 客户端超时={TIMEOUT:.0f}s", flush=True)

# ── 第 1 轮提示词（label_prompt_round1.md §A）──────────────────────────────
# 除开头与【输出格式】两段外，与恢复出的原文逐字一致。
SYSTEM_PROMPT = """你是数据清洗标注员，协助传播学研究。下面给你一批帖子（JSONL 格式；字段：row_id, platform, author, title, match_sentence，部分行另有 ctx_full_content 补充上下文）。

【研究背景】数据为 2026 年 1-5 月社交媒体（小红书/抖音/微博）上关键词"张雪峰"的舆情抓取结果。张雪峰（教育博主，峰学蔚来创始人）于 2026 年 3 月 24 日去世，数据中大量涉及该事件及其衍生网络梗（雪碧梗、巧乐兹梗、"牢张"、捐五千万、天地银行、"上帝想买提分笔记"、抽象文化等）。研究课题是张雪峰相关"梗"文化。爬虫把蹭"张雪峰"名字做推广的广告也抓了进来，现在要标注每行的有效性。

【判定核心标准】判断"张雪峰在句中扮演什么角色"：他是讨论对象（有效），还是被借来引流的招牌（无效广告），还是内容根本不是真实帖子（噪音）。辅助判断法：把"张雪峰"三个字拿掉，剩下的内容若是商品/机构推销，就是借势营销；剩下的仍是关于他的讨论或玩梗，就是有效。

【类别体系】category 只能取以下枚举值：
无效类：
- 借势营销：真实目的是推广课程、资料、题库、押题、留学/国际本科中介、志愿填报付费服务、招聘、带货、App、加群加微信、关注引流等；"张雪峰建议/力荐/说得对"只是权威背书。特征：购买/报名/领取/咨询引导、价格学费、机构名、联系方式。
- 爬取噪音：非真实用户帖子——页面 UI 文本（"小红书号：""IP属地：""粉丝获赞""橱窗逛逛""直播动态"）、AI 对话残留（"ai总结""很抱歉，我可能误解了你的问题"）、题库/网课界面文本（"第1章""下载试卷""正确率0%"）、与张雪峰无实质关联的碎片或纯标签堆砌。
有效类：
- 梗文化讨论：玩梗、二创、黑色幽默、抽象话。两个边界规则：①梗帖里提到"买课/提分笔记/一个亿/天地银行/提分"通常是玩梗（如"上帝想买提分笔记所以带走了牢张""花一个亿买你的课，通过天地银行支付"），不是广告；②只有话题标签、无其他文字的纯梗标签帖（如"#张雪峰 #巧乐兹 #雪碧"）归此类——标签本身即参与行为。
- 事件悼念讨论：关于他离世、告别仪式、生前事迹的严肃讨论、追思、"一路走好"类内容。
- 教育观点讨论：以他或其言论为讨论主体的正经讨论——志愿填报、考研、就业观点的引用与评价、支持或批评。语录摘编/干货搬运按引流意图划线：无推销话术且无引流引导（关注/私信/主页/领资料）→归此类；机构/资料/规划/学姐类账号发布，或含引流引导→归借势营销。
- 其他讨论：真实用户内容、张雪峰确为讨论主体，但不属于上述各类。
存疑类：
- 存疑：仅当信息确实不足（截断且无 ctx_full_content 可依、语义完全不明）才使用，预期占比≤2%，不要滥用；有 ctx_full_content 的行优先用它判定。

【输出格式 - 严格 JSON 数组，不要任何其他文字】
对每条帖子依次判定，输出一个 JSON 数组，元素个数和顺序与输入帖子一一对应：
[{"row_id": 整数, "valid": "有效|无效|存疑", "category": 枚举值, "reason": "≤30字判定依据"}, ...]
valid 与 category 必须一致（无效→借势营销/爬取噪音；有效→梗文化讨论/事件悼念讨论/教育观点讨论/其他讨论；存疑→存疑）。
不得遗漏行。"""

VALID_CATS = {"借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音", "存疑"}
VALID_MAP = {
    "有效": {"梗文化讨论", "事件悼念讨论", "教育观点讨论", "其他讨论"},
    "无效": {"借势营销", "爬取噪音"},
    "存疑": {"存疑"},
}

# 源表列索引，main() 读取表头后填充
IDX = {}


def s(v):
    """复刻第 1 轮构建代码里的 str(df.at[i, col])：pandas 空值(NaN) → 'nan'。"""
    return "nan" if v is None else str(v)


def build_record(rid, row):
    """逐字复刻第 1 轮批次构建：/tmp/zxf_full/batch_*.jsonl 的每行。"""
    ms = re.sub(r"</?em>", "", "" if row[IDX["match_sentence"]] is None
                else str(row[IDX["match_sentence"]]))[:400]
    rec = {
        "row_id": rid,
        "platform": s(row[IDX["media_type"]]),
        "author": s(row[IDX["author_name"]])[:30],
        "title": s(row[IDX["title"]])[:80],
        "match_sentence": ms,
    }
    if len(ms) < 15:  # 截断/空的行，补上下文供判定
        fc = s(row[IDX["full_content"]])[:250] if row[IDX["full_content"]] is not None else ""
        rec["ctx_full_content"] = fc if fc else s(row[IDX["content"]])[:250]
    return rec


def extract_array(txt, expect):
    """提取长度为 expect 的 JSON 数组，按 row_id 对齐并做值域校验。"""
    if txt.startswith("```"):
        txt = txt.split("```")[1] if "```" in txt else txt
        if txt.startswith("json"):
            txt = txt[4:]
        txt = txt.strip()
    lo, hi = txt.find("["), txt.rfind("]")
    if lo < 0 or hi <= lo:
        raise ValueError(f"no json array: {txt[:60]}")
    arr = json.loads(txt[lo:hi + 1])
    if not isinstance(arr, list) or len(arr) != expect:
        raise ValueError(f"array len {len(arr) if isinstance(arr, list) else '?'} != {expect}")
    out = {}
    for obj in arr:
        rid = obj.get("row_id")
        if not isinstance(rid, int) or rid in out:
            raise ValueError(f"bad row_id: {rid}")
        cat = obj.get("category", "其他讨论")
        valid = obj.get("valid", "有效")
        reason = str(obj.get("reason", ""))[:200]
        if cat not in VALID_CATS:
            cat = "其他讨论"
        if valid not in VALID_MAP:
            valid = "有效"
        if cat not in VALID_MAP[valid]:      # 第1轮规则：valid 与 category 必须一致
            valid = next(v for v, cs in VALID_MAP.items() if cat in cs)
        out[rid] = (cat, reason, valid)
    return out


def label_batch(row_ids, records):
    """records=[rec,...]（第1轮格式）。返回 [{row_num,category,reason,valid,failed}]。"""
    n = len(records)
    user_msg = "请依次判定以下 %d 条帖子：\n\n" % n + "\n".join(
        json.dumps(r, ensure_ascii=False) for r in records)
    for attempt in range(3):
        try:
            r = CLIENT.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_msg},
                ],
                temperature=0.1, max_tokens=MAX_TOKENS,
            )
            txt = (r.choices[0].message.content or "").strip()
            if not txt:
                raise ValueError("empty content (thinking exhausted max_tokens?)")
            got = extract_array(txt, n)
            recs = []
            for rid in row_ids:
                if rid not in got:
                    raise ValueError(f"missing row_id {rid}")
                cat, reason, valid = got[rid]
                recs.append({"row_num": rid + 2, "category": cat, "reason": reason,
                             "valid": valid, "failed": False})
            return recs
        except Exception as e:
            if attempt == 2:
                return [{"row_num": rid + 2, "category": "其他讨论",
                         "reason": f"[API失败:{str(e)[:80]}]", "valid": "无效", "failed": True}
                        for rid in row_ids]
            time.sleep(2 ** attempt)
    return [{"row_num": rid + 2, "category": "其他讨论", "reason": "[重试失败]",
             "valid": "无效", "failed": True} for rid in row_ids]


def load_done():
    done = {}
    if os.path.exists(JSONL):
        with open(JSONL) as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    o = json.loads(line)
                    done[o["row_num"]] = o
                except Exception:
                    continue
    return done


def append_jsonl(recs):
    with open(JSONL, "a") as f:
        for rec in recs:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def run_pool(batch_jobs, done_rows, total_rows, t0):
    buf = []
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = [ex.submit(label_batch, rids, recs) for rids, recs in batch_jobs]
        for fut in as_completed(futs):
            recs = fut.result()
            buf.extend(recs)
            done_rows += len(recs)
            if len(buf) >= 100:
                append_jsonl(buf)
                buf = []
            if done_rows % 200 < BSIZE:
                elapsed = time.time() - t0
                rate = done_rows / elapsed if elapsed > 0 else 0
                remaining = (total_rows - done_rows) / rate if rate > 0 else 0
                print(f"进度 {done_rows}/{total_rows} ({done_rows*100//total_rows}%) | "
                      f"速度 {rate:.1f} 条/秒 | 已跑 {elapsed/60:.1f} 分钟 | "
                      f"预计剩余 {remaining/60:.1f} 分钟", flush=True)
                with open(PROGRESS, "w") as f:
                    json.dump({"done_count": done_rows, "total": total_rows,
                               "model": MODEL, "slug": SLUG, "mode": "batch",
                               "prompt": "round1_recovered"}, f)
    if buf:
        append_jsonl(buf)
    return done_rows


def main():
    t0 = time.time()
    print("读取源数据...", flush=True)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    for name in ["title", "content", "full_content", "media_type", "author_name", "match_sentence"]:
        IDX[name] = header.index(name)
    rows, recs = {}, {}
    rn = 1
    for row in it:
        rn += 1
        if row is None or all(v is None for v in row):
            continue
        rows[rn] = row
        recs[rn] = build_record(rn - 2, row)     # row_id = 0 基位置 = Excel行号 - 2
    wb.close()
    total = len(rows)
    print(f"共 {total} 条源数据（row_id 0–{total-1}）", flush=True)
    nctx = sum(1 for r in recs.values() if "ctx_full_content" in r)
    print(f"输入字段校验：补 ctx_full_content 的行数 = {nctx}（第1轮记录 1312）", flush=True)

    done = load_done()
    ok = {n for n, o in done.items() if not o.get("failed")}
    print(f"断点续跑：已完成 {len(ok)} 条，其中失败待重试 {len(done) - len(ok)} 条", flush=True)
    with open(PROGRESS, "w") as f:
        json.dump({"done_count": len(ok), "total": total, "model": MODEL, "slug": SLUG,
                   "mode": "batch", "prompt": "round1_recovered"}, f)

    pending = [n for n in sorted(rows) if n not in ok]
    batch_jobs = []
    for i in range(0, len(pending), BSIZE):
        chunk = pending[i:i + BSIZE]
        batch_jobs.append(([n - 2 for n in chunk], [recs[n] for n in chunk]))
    print(f"待处理：{len(pending)} 条，分为 {len(batch_jobs)} 批（每批 {BSIZE} 条）", flush=True)

    # 一次性提交全部批次。分块（原为 WORKERS*3）会在每块末尾等最慢的那次调用，
    # 而思考型模型的延迟长尾很长（实测单次 110–430s），块边界会整池空等：
    # 6405 批 / 192 = 33 个边界，累计可达数小时。单池则工作线程不空转。
    done_rows = run_pool(batch_jobs, 0, len(pending), t0)

    # 失败重试（最多 2 轮）
    for rnd in range(1, 3):
        cur = load_done()
        failed = [n for n in sorted(rows) if n in cur and cur[n].get("failed")]
        if not failed:
            break
        print(f"失败重试第 {rnd} 轮：{len(failed)} 条", flush=True)
        rj = []
        for i in range(0, len(failed), BSIZE):
            chunk = failed[i:i + BSIZE]
            rj.append(([n - 2 for n in chunk], [recs[n] for n in chunk]))
        run_pool(rj, 0, len(failed), t0)

    # 汇总装配
    print("装配输出 Excel...", flush=True)
    labels = load_done()
    missing = [n for n in rows if n not in labels or labels[n].get("failed")]
    if missing:
        print(f"❌ 有 {len(missing)} 条仍无有效判定，中止装配。前5条: {missing[:5]}", flush=True)
        sys.exit(2)

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.append(header)
    cat_c, val_c, rl = Counter(), Counter(), []
    for n in sorted(rows):
        lab = labels[n]
        out_ws.append([lab["valid"], lab["category"], lab["reason"]] + list(rows[n])[3:])
        cat_c[lab["category"]] += 1
        val_c[lab["valid"]] += 1
        rl.append(len(lab["reason"]))
    out_wb.save(OUT)

    rl.sort()
    elapsed = time.time() - t0
    print(f"\n✅ 全部完成！总耗时 {elapsed/60:.1f} 分钟，输出：{OUT}", flush=True)
    print(f"类别分布：{dict(cat_c.most_common())}", flush=True)
    print(f"有效性分布：{dict(val_c)}", flush=True)
    print(f"判定依据长度：中位 {rl[len(rl)//2]}  最大 {max(rl)}  "
          f"≤30字 {sum(1 for x in rl if x <= 30)/len(rl)*100:.2f}%", flush=True)


if __name__ == "__main__":
    main()
