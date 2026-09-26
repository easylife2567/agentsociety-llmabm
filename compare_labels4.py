#!/usr/bin/env python3
"""
四模型标签一致性对比：DeepSeek-V4 原标 / doubao-seed / GLM / DeepSeek-V4.1
按行对齐读取四份 xlsx 的前三列（数据有效性/内容类别/判定依据），输出：
  - 各评审方类别/有效性分布
  - 两两总体一致率 + Cohen's kappa（类别、有效性）+ 混淆矩阵 + 分歧示例
  - 四方同时一致率、Fleiss' kappa（多评审者信度）
  - 多数投票（严格多数 ≥3/4）分布、各评审方与多数投票的一致率、无多数行清单
  - DeepSeek 原标「存疑」行专列
用法: python compare_labels4.py
"""
import os, sys, datetime
from collections import Counter
from itertools import combinations
import openpyxl
from openpyxl import Workbook

ROOT = "/Users/easylife/Project/AgentSociety"
CAT_ORDER = ["借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音", "存疑"]
VAL_ORDER = ["有效", "无效", "存疑"]

PATHS = {
    "DeepSeek-V4": f"{ROOT}/抖音微博小红书-全量已打标.xlsx",
    "doubao-seed": f"{ROOT}/抖音微博小红书-独立重打标_Seed2_1_lite.xlsx",
    "GLM": f"{ROOT}/抖音微博小红书-独立重打标_GLM5_3flash.xlsx",
    "DeepSeek-V4.1": f"{ROOT}/抖音微博小红书-独立重打标_DSv4_1_flash.xlsx",
}


def load_labels(path):
    """返回 [(valid, cat, reason, rid, snippet)]，按行序。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    id_i, t_i, c_i = header.index("id"), header.index("title"), header.index("content")
    out = []
    for row in it:
        valid, cat, reason = str(row[0]), str(row[1]), str(row[2])
        rid = row[id_i]
        title = row[t_i] or ""
        content = row[c_i] or ""
        snippet = str(title)[:40] if str(title).strip() else str(content)[:40]
        out.append((valid, cat, reason, rid, snippet))
    wb.close()
    return out


def kappa(a, b):
    n = len(a)
    if n == 0:
        return float("nan")
    joint = Counter(zip(a, b))
    p_o = sum(v for (x, y), v in joint.items() if x == y) / n
    ca, cb = Counter(a), Counter(b)
    p_e = sum(ca[k] * cb.get(k, 0) for k in ca) / (n * n)
    if p_e == 1:
        return 1.0
    return (p_o - p_e) / (1 - p_e)


def fleiss_kappa(labels_by_rater):
    """labels_by_rater: {rater: [label,...]}，全部评审方覆盖全部样本（无缺失）。"""
    raters = list(labels_by_rater)
    n = len(labels_by_rater[raters[0]])
    m = len(raters)
    if n == 0 or m < 2:
        return float("nan")
    all_cats = sorted({l for r in raters for l in labels_by_rater[r]})
    p_i_sum = 0.0
    cat_tot = Counter()
    for i in range(n):
        cnt = Counter(labels_by_rater[r][i] for r in raters)
        cat_tot.update(cnt)
        p_i_sum += (sum(c * c for c in cnt.values()) - m) / (m * (m - 1))
    p_bar = p_i_sum / n
    p_e = sum((cat_tot[c] / (n * m)) ** 2 for c in all_cats)
    if p_e == 1:
        return 1.0
    return (p_bar - p_e) / (1 - p_e)


def confusion(a, b, order):
    m = {x: {y: 0 for y in order} for x in order}
    for x, y in zip(a, b):
        m.setdefault(x, {y: 0 for y in order})
        m[x].setdefault(y, 0)
        m[x][y] += 1
    return m


def fmt_matrix(m, order):
    short = {c: c[:2] for c in order}
    head = "A\\B | " + " | ".join(f"{short[c]:>4}" for c in order)
    lines = [head, "-" * len(head)]
    for x in order:
        lines.append(f"{short[x]:>4} | " + " | ".join(f"{m.get(x, {}).get(y, 0):>4}" for y in order))
    return "\n".join(lines)


def main():
    data = {}
    for name, p in PATHS.items():
        if not os.path.exists(p):
            print(f"❌ 缺少评审方文件: {name} ({p})")
            sys.exit(1)
        data[name] = load_labels(p)
        print(f"载入 {name}: {len(data[name])} 行")

    names = list(data)
    n = len(data[names[0]])
    for nm in names:
        assert len(data[nm]) == n, f"{nm} 行数不一致: {len(data[nm])} != {n}"
    for i in [0, n // 2, n - 1]:
        ids = {str(data[nm][i][3]) for nm in names}
        assert len(ids) == 1, f"行 {i+2} id 不对齐: {ids}"
    print(f"行数与 id 对齐校验通过（n={n}，评审方={len(names)}）")

    cats = {nm: [r[1] for r in data[nm]] for nm in names}
    vals = {nm: [r[0] for r in data[nm]] for nm in names}

    today = datetime.date.today().strftime("%Y%m%d")
    md = [f"# 四模型标签一致性对比报告（{today}）\n"]
    md.append(f"- 评审方（4）：{'、'.join(names)}")
    md.append(f"- 样本量：{n} 条")
    md.append("- 三轮独立重打标均未读取任何已有标签列。")
    md.append("- **编码配置并非四轮一致（2026-09-26 复核更正）**：后三轮（doubao-seed / GLM / DeepSeek-V4.1）"
              "所用判定标准（六类定义、判定原则）逐字相同，仅输出格式段与每次调用的帖文条数不同"
              "（doubao-seed 逐条；GLM 与 DeepSeek-V4.1 每 8 条一次，GLM 轮前 6.3% 的行为逐条判定）；"
              "**第一轮（DeepSeek-V4 原标）所用提示词与之不同且未留存**——其「判定依据」中位长度 11 字"
              "（99.5% ≤20 字），后三轮为 27—34 字（≤20 字仅 3.6%—12.7%），输出端签名呈类别式差异。"
              "故原标参与的三对 κ 同时含口径差异，只宜作参考上界；后三轮之间的 κ 才是口径可比的部分。"
              "复核脚本：`prompt_provenance_check.py`。")
    md.append("")
    md.append("**数据质量说明**：")
    md.append("- GLM 轮 184 条（0.35%）因方舟内容审核拒绝/模型拒答/无有效输出，标记为 其他讨论+无效（判定依据含「需人工复核」）。")
    md.append("- DeepSeek-V4.1 轮 9 条（0.017%）经批量连坐隔离后的两轮单条重试仍被模型稳定拒答（内容本身为普通帖，属模型侧过度拒答），"
              "按同一约定标记为 其他讨论+无效（判定依据含「需人工复核」）。")
    md.append("- doubao-seed 轮成品已验证与 v1 轮输出【新】三列 100% 同源。")
    md.append("")

    for nm in names:
        cc, vv = Counter(cats[nm]), Counter(vals[nm])
        md.append(f"- **{nm}** 类别分布：{'，'.join(f'{k} {v} ({v*100/n:.1f}%)' for k, v in cc.most_common())}")
        md.append(f"- **{nm}** 有效性分布：{'，'.join(f'{k} {v}' for k, v in vv.most_common())}")
    md.append("")

    xwb = Workbook()
    xwb.remove(xwb.active)

    # ---- 两两对比 ----
    for a, b in combinations(names, 2):
        agree_c = sum(x == y for x, y in zip(cats[a], cats[b]))
        agree_v = sum(x == y for x, y in zip(vals[a], vals[b]))
        k_c = kappa(cats[a], cats[b])
        k_v = kappa(vals[a], vals[b])
        order_c = [c for c in CAT_ORDER if c in set(cats[a]) | set(cats[b])]
        order_v = [v for v in VAL_ORDER if v in set(vals[a]) | set(vals[b])]
        md.append(f"## {a} vs {b}\n")
        md.append(f"- 类别一致率：{agree_c}/{n} = {agree_c*100/n:.2f}%，Cohen's κ = {k_c:.4f}")
        md.append(f"- 有效性一致率：{agree_v}/{n} = {agree_v*100/n:.2f}%，Cohen's κ = {k_v:.4f}")
        md.append(f"\n类别混淆矩阵（行={a}，列={b}）：\n```\n{fmt_matrix(confusion(cats[a], cats[b], order_c), order_c)}\n```")
        md.append(f"\n有效性混淆矩阵：\n```\n{fmt_matrix(confusion(vals[a], vals[b], order_v), order_v)}\n```")

        dis = [(i, data[a][i], data[b][i]) for i in range(n) if cats[a][i] != cats[b][i]]
        md.append(f"- 类别分歧行数：{len(dis)}")
        md.append("\n分歧示例（行号/内容摘要/两方判定）：\n")
        for i, ra, rb in dis[:20]:
            md.append(f"- 行{i+2}｜{ra[4]}｜{a}: {ra[1]}（{ra[0]}）←→ {b}: {rb[1]}（{rb[0]}）")
        md.append("")

        ws = xwb.create_sheet(f"{a}-vs-{b}-混淆")
        ws.append(["A\\B"] + order_c)
        cm = confusion(cats[a], cats[b], order_c)
        for x in order_c:
            ws.append([x] + [cm.get(x, {}).get(y, 0) for y in order_c])
        wsd = xwb.create_sheet(f"{a}-vs-{b}-分歧")
        wsd.append(["行号", "id", "内容摘要", f"{a}类别", f"{a}有效性", f"{b}类别", f"{b}有效性", f"{a}依据", f"{b}依据"])
        for i, ra, rb in dis:
            wsd.append([i + 2, str(ra[3]), ra[4], ra[1], ra[0], rb[1], rb[0], ra[2], rb[2]])

    # ---- 多方同时一致 + Fleiss' κ ----
    agree_all_c = sum(len({cats[nm][i] for nm in names}) == 1 for i in range(n))
    agree_all_v = sum(len({vals[nm][i] for nm in names}) == 1 for i in range(n))
    fk_c = fleiss_kappa(cats)
    fk_v = fleiss_kappa(vals)
    md.append(f"## 四方同时一致与多评审者信度\n")
    md.append(f"- 类别四方一致：{agree_all_c}/{n} = {agree_all_c*100/n:.2f}%")
    md.append(f"- 有效性四方一致：{agree_all_v}/{n} = {agree_all_v*100/n:.2f}%")
    md.append(f"- Fleiss' κ（类别，4 评审者）：{fk_c:.4f}")
    md.append(f"- Fleiss' κ（有效性，4 评审者）：{fk_v:.4f}\n")

    # ---- 多数投票（严格多数 ≥3/4） ----
    def majority(seq):
        c = Counter(seq)
        top, cnt = c.most_common(1)[0]
        return top if cnt >= 3 else None

    maj_c = [majority([cats[nm][i] for nm in names]) for i in range(n)]
    maj_v = [majority([vals[nm][i] for nm in names]) for i in range(n)]
    no_maj_c = sum(1 for x in maj_c if x is None)
    no_maj_v = sum(1 for x in maj_v if x is None)
    md.append("## 多数投票（严格多数 ≥3/4）\n")
    cc = Counter(x for x in maj_c if x)
    vv = Counter(x for x in maj_v if x)
    md.append(f"- 类别有严格多数行：{n - no_maj_c}/{n} = {(n-no_maj_c)*100/n:.2f}%（无多数 {no_maj_c} 条）")
    md.append(f"- 多数类别分布：{'，'.join(f'{k} {v} ({v*100/n:.1f}%)' for k, v in cc.most_common())}")
    md.append(f"- 有效性有严格多数行：{n - no_maj_v}/{n} = {(n-no_maj_v)*100/n:.2f}%（无多数 {no_maj_v} 条）")
    md.append(f"- 多数有效性分布：{'，'.join(f'{k} {v}' for k, v in vv.most_common())}")
    md.append("\n各评审方与多数投票的一致率（仅统计有严格多数的行）：\n")
    for nm in names:
        denom_c = n - no_maj_c
        denom_v = n - no_maj_v
        ac = sum(1 for i in range(n) if maj_c[i] and cats[nm][i] == maj_c[i]) / denom_c
        av = sum(1 for i in range(n) if maj_v[i] and vals[nm][i] == maj_v[i]) / denom_v
        md.append(f"- {nm}：类别 {ac*100:.2f}%，有效性 {av*100:.2f}%")
    md.append("")

    wsm = xwb.create_sheet("多数投票")
    wsm.append(["行号", "id", "内容摘要", "多数类别", "多数有效性"] + [f"{nm}类别/{nm}有效性" for nm in names])
    for i in range(n):
        row = [i + 2, str(data[names[0]][i][3]), data[names[0]][i][4],
               maj_c[i] or "无多数", maj_v[i] or "无多数"]
        for nm in names:
            row += [cats[nm][i], vals[nm][i]]
        wsm.append(row)

    # ---- 四方分歧明细 ----
    ws4 = xwb.create_sheet("四方分歧")
    ws4.append(["行号", "id", "内容摘要"] + [f"{nm}类别/{nm}有效性" for nm in names])
    n_dis4 = 0
    for i in range(n):
        if len({cats[nm][i] for nm in names}) > 1:
            n_dis4 += 1
            row = [i + 2, str(data[names[0]][i][3]), data[names[0]][i][4]]
            for nm in names:
                row += [cats[nm][i], vals[nm][i]]
            ws4.append(row)
    md.append(f"## 类别四方分歧\n\n- 分歧行数：{n_dis4}/{n} = {n_dis4*100/n:.2f}%（明细见 xlsx「四方分歧」表）\n")

    # ---- DeepSeek 原标「存疑」行 ----
    dq = [i for i in range(n) if cats["DeepSeek-V4"][i] == "存疑"]
    md.append(f"## DeepSeek-V4 原标「存疑」行（{len(dq)} 条）\n")
    for i in dq:
        others = "，".join(f"{nm}: {cats[nm][i]}（{vals[nm][i]}）" for nm in names if nm != "DeepSeek-V4")
        md.append(f"- 行{i+2}｜{data['DeepSeek-V4'][i][4]}｜{others}")
    wsq = xwb.create_sheet("DeepSeek存疑行")
    wsq.append(["行号", "id", "内容摘要", "DeepSeek-V4类别"] + [f"{nm}类别" for nm in names if nm != "DeepSeek-V4"])
    for i in dq:
        wsq.append([i + 2, str(data["DeepSeek-V4"][i][3]), data["DeepSeek-V4"][i][4], "存疑"] +
                   [cats[nm][i] for nm in names if nm != "DeepSeek-V4"])

    md_path = f"{ROOT}/四模型标签一致性对比报告_{today}.md"
    xlsx_path = f"{ROOT}/四模型标签一致性对比明细_{today}.xlsx"
    with open(md_path, "w") as f:
        f.write("\n".join(md))
    xwb.save(xlsx_path)
    print(f"\n✅ 报告：{md_path}\n✅ 明细：{xlsx_path}")


if __name__ == "__main__":
    main()
