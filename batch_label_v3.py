#!/usr/bin/env python3
"""
独立批量打标脚本 v3（GLM 轮·批量版）- 每次 API 调用判定 8 条帖子
起因：glm-5-3-flash 为强制思考型模型，逐条调用等效延时 ~25s，24 并发仅 ~1 条/秒。
批量后单次调用产出 8 条判定，吞吐提升数倍。判定标准（六类定义/判定原则）与前两轮完全一致。

用法: python batch_label_v3.py <model_id> <slug> [workers] [batch_size]
产出: 抖音微博小红书-独立重打标_<slug>.xlsx（结构=原表57列，前三列替换为新判定）
断点: .label_results_<slug>.jsonl（与 v2 逐条版同 schema，按行号续跑、互相兼容）
"""
import os, sys, json, time, dotenv, openpyxl
from openpyxl import Workbook
from openai import OpenAI
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import Counter

dotenv.load_dotenv("/Users/easylife/Project/AgentSociety/.env")
API_KEY = os.environ["AGENTSOCIETY_LLM_API_KEY"]
API_BASE = os.environ["AGENTSOCIETY_LLM_API_BASE"]

MODEL = sys.argv[1]
SLUG = sys.argv[2]
WORKERS = int(sys.argv[3]) if len(sys.argv) > 3 else 24
BSIZE = int(sys.argv[4]) if len(sys.argv) > 4 else 8
MAX_TOKENS = int(os.environ.get("LABEL_MAX_TOKENS", "8000"))

ROOT = "/Users/easylife/Project/AgentSociety"
SRC = f"{ROOT}/抖音微博小红书-全量已打标.xlsx"
OUT = f"{ROOT}/抖音微博小红书-独立重打标_{SLUG}.xlsx"
JSONL = f"{ROOT}/.label_results_{SLUG}.jsonl"
PROGRESS = f"{ROOT}/.label_progress_{SLUG}.json"

print(f"[配置] 模型={MODEL} | slug={SLUG} | 并发={WORKERS} | 每次调用判定 {BSIZE} 条 | max_tokens={MAX_TOKENS}", flush=True)
print(f"[配置] 输出={OUT}", flush=True)

CLIENT = OpenAI(api_key=API_KEY, base_url=API_BASE, timeout=180)

# 判定标准与 batch_label.py / batch_label_v2.py 完全一致；仅输出格式改为数组
SYSTEM_PROMPT = """你是一个中文社交媒体内容分析助手。请对每条帖子进行三方面判定。

【内容类别 - 六选一，不允许自创】
- 借势营销：以商品、服务、账号引流、带货、二维码、价格、链接、限时优惠等为目的的推广内容；或借某热点人物/事件为特定产品/服务背书引流
- 事件悼念讨论：对逝者、事故、灾害、公共事件的缅怀、哀悼、点烛、RIP、默哀类内容；或围绕此类事件的严肃讨论
- 教育观点讨论：知识科普、观点输出、教育理念讨论、人生感悟、有信息增量的分析类内容（非营销）
- 梗文化讨论：meme、谐音梗、段子、玩梗、跟风玩梗、无实义的网络流行语、搞笑吐槽类内容
- 其他讨论：以上四类都不符合，但仍是正常有意义的帖子（如日常分享、生活记录、提问求助、情感抒发）
- 爬取噪音：乱码、纯表情、@好友无意义内容、无意义短文本、广告刷屏残留、空内容、格式损坏等无实质信息的条目

【数据有效性 - 二选一】
- 有效：包含可用于后续研究分析的实质内容
- 无效：无实质内容、纯噪音、纯营销引流无研究价值、或内容过于残缺无法分析

【判定原则】
1. 先判断是否为噪音：内容残缺/乱码/纯@人/纯表情 → 爬取噪音，无效
2. 再判断是否为营销：带货/引流/推广 → 借势营销，无效
3. 再判断是否为悼念：对逝者/事故/灾害的缅怀 → 事件悼念讨论，有效
4. 再判断是否为玩梗：meme/段子/谐音梗/跟风玩梗 → 梗文化讨论，有效
5. 再判断是否为教育观点：知识输出/观点分析/教育理念/人生感悟 → 教育观点讨论，有效
6. 以上都不符合 → 其他讨论，根据是否有实质内容判断有效性

【输出格式 - 严格 JSON 数组，不要任何其他文字】
对用户消息里的每条帖子依次判定，输出一个 JSON 数组，元素个数和顺序与输入帖子一一对应：
[{"no": 1, "category": "类别名", "reason": "一句话依据", "valid": "有效|无效"}, ...]"""

VALID_CATS = {"借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音"}

# 源表列索引，main() 读取表头后填充，run_pool() 使用
IDX = {}


def fmt_post(k, title, content, media):
    return f"【帖子{k}】\n平台：{media or ''}\n标题：{title or ''}\n正文：{(content or '')[:2000]}"


def extract_array(txt, expect):
    """从回复中提取长度为 expect 的 JSON 数组并做值域校验，返回 {no: (cat, reason, valid)}。"""
    if txt.startswith("```"):
        txt = txt.split("```")[1] if "```" in txt else txt
        if txt.startswith("json"):
            txt = txt[4:]
        txt = txt.strip()
    lo, hi = txt.find("["), txt.rfind("]")
    if lo < 0 or hi <= lo:
        raise ValueError(f"no json array: {txt[:60]}")
    arr = json.loads(txt[lo:hi + 1])
    if not isinstance(arr, list) or len(arr) != expect:
        raise ValueError(f"array len {len(arr) if isinstance(arr, list) else '?'} != {expect}")
    out = {}
    for obj in arr:
        no = obj.get("no")
        if not isinstance(no, int) or no in out or not (1 <= no <= expect):
            raise ValueError(f"bad no: {no}")
        cat = obj.get("category", "其他讨论")
        reason = str(obj.get("reason", ""))[:200]
        valid = obj.get("valid", "有效")
        if cat not in VALID_CATS:
            cat = "其他讨论"
        if valid not in {"有效", "无效"}:
            valid = "有效"
        out[no] = (cat, reason, valid)
    return out


def label_batch(row_nums, posts):
    """posts=[(title,content,media)]，返回 [{"row_num","category","reason","valid","failed"}]。
    注意：必须返回 dict 而非元组——append_jsonl 按字典序列化，元组会被写成数组行导致断点解析失败。"""
    n = len(posts)
    user_msg = f"请依次判定以下 {n} 条帖子：\n\n" + "\n\n".join(
        fmt_post(i + 1, *posts[i]) for i in range(n))
    for attempt in range(3):
        try:
            r = CLIENT.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_msg},
                ],
                temperature=0.1, max_tokens=MAX_TOKENS,
            )
            txt = (r.choices[0].message.content or "").strip()
            if not txt:
                raise ValueError("empty content (thinking exhausted max_tokens?)")
            got = extract_array(txt, n)
            recs = []
            for k, rn in enumerate(row_nums):
                cat, reason, valid = got[k + 1]
                recs.append({"row_num": rn, "category": cat, "reason": reason,
                             "valid": valid, "failed": False})
            return recs
        except Exception as e:
            if attempt == 2:
                return [{"row_num": rn, "category": "其他讨论",
                         "reason": f"[API失败:{str(e)[:80]}]", "valid": "无效", "failed": True}
                        for rn in row_nums]
            time.sleep(2 ** attempt)
    return [{"row_num": rn, "category": "其他讨论", "reason": "[重试失败]",
             "valid": "无效", "failed": True} for rn in row_nums]


def load_done():
    done = {}
    if os.path.exists(JSONL):
        with open(JSONL) as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    o = json.loads(line)
                    done[o["row_num"]] = o
                except Exception:
                    continue
    return done


def append_jsonl(recs):
    with open(JSONL, "a") as f:
        for rec in recs:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def run_pool(batch_jobs, done_rows, total_rows, t0):
    """batch_jobs=[(row_nums, posts)]，结果按行追加 JSONL。"""
    buf = []
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = [ex.submit(label_batch, rns, ps) for rns, ps in batch_jobs]
        for fut in as_completed(futs):
            recs = fut.result()
            buf.extend(recs)
            done_rows += len(recs)
            if len(buf) >= 100:
                append_jsonl(buf)
                buf = []
            if done_rows % 200 < BSIZE:
                elapsed = time.time() - t0
                rate = done_rows / elapsed if elapsed > 0 else 0
                remaining = (total_rows - done_rows) / rate if rate > 0 else 0
                print(f"进度 {done_rows}/{total_rows} ({done_rows*100//total_rows}%) | "
                      f"速度 {rate:.1f} 条/秒 | 已跑 {elapsed/60:.1f} 分钟 | "
                      f"预计剩余 {remaining/60:.1f} 分钟", flush=True)
                with open(PROGRESS, "w") as f:
                    json.dump({"done_count": done_rows, "total": total_rows,
                               "model": MODEL, "slug": SLUG, "mode": "batch"}, f)
    if buf:
        append_jsonl(buf)
    return done_rows


def main():
    t0 = time.time()
    print("读取源数据...", flush=True)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    IDX.update({name: header.index(name) for name in ["title", "content", "full_content", "media_type"]})
    rows = {}
    rn = 1
    for row in it:
        rn += 1
        if row is None or all(v is None for v in row):
            continue
        rows[rn] = row
    wb.close()
    total = len(rows)
    print(f"共 {total} 条源数据", flush=True)

    done = load_done()
    ok = {n for n, o in done.items() if not o.get("failed")}
    print(f"断点续跑：已完成 {len(ok)} 条，其中失败待重试 {len(done) - len(ok)} 条", flush=True)
    with open(PROGRESS, "w") as f:
        json.dump({"done_count": len(ok), "total": total, "model": MODEL, "slug": SLUG, "mode": "batch"}, f)

    pending = [n for n in sorted(rows) if n not in ok]

    # 组批：相邻行号连续成批
    batch_jobs = []
    for i in range(0, len(pending), BSIZE):
        chunk = pending[i:i + BSIZE]
        posts = [(rows[n][IDX["title"]], rows[n][IDX["content"]] or rows[n][IDX["full_content"]],
                  rows[n][IDX["media_type"]]) for n in chunk]
        batch_jobs.append((chunk, posts))
    print(f"待处理：{len(pending)} 条，分为 {len(batch_jobs)} 批（每批 {BSIZE} 条）", flush=True)

    done_rows = 0
    CHUNK = WORKERS * 3
    for i in range(0, len(batch_jobs), CHUNK):
        done_rows = run_pool(batch_jobs[i:i + CHUNK], done_rows, len(pending), t0)

    # 失败重试（最多 2 轮）
    for rnd in range(1, 3):
        cur = load_done()
        failed = [n for n in sorted(rows) if n in cur and cur[n].get("failed")]
        if not failed:
            break
        print(f"失败重试第 {rnd} 轮：{len(failed)} 条", flush=True)
        rj = []
        for i in range(0, len(failed), BSIZE):
            chunk = failed[i:i + BSIZE]
            posts = [(rows[n][IDX["title"]], rows[n][IDX["content"]] or rows[n][IDX["full_content"]],
                      rows[n][IDX["media_type"]]) for n in chunk]
            rj.append((chunk, posts))
        run_pool(rj, 0, len(failed), t0)

    # 汇总装配
    print("装配输出 Excel...", flush=True)
    labels = load_done()
    missing = [n for n in rows if n not in labels or labels[n].get("failed")]
    if missing:
        print(f"❌ 有 {len(missing)} 条仍无有效判定，中止装配。前5条: {missing[:5]}", flush=True)
        sys.exit(2)

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.append(header)  # 表头与原表完全一致
    cat_c, val_c = Counter(), Counter()
    for n in sorted(rows):
        row = rows[n]
        lab = labels[n]
        out_ws.append([lab["valid"], lab["category"], lab["reason"]] + list(row)[3:])
        cat_c[lab["category"]] += 1
        val_c[lab["valid"]] += 1
    out_wb.save(OUT)

    elapsed = time.time() - t0
    print(f"\n✅ 全部完成！总耗时 {elapsed/60:.1f} 分钟，输出：{OUT}", flush=True)
    print(f"类别分布：{dict(cat_c.most_common())}", flush=True)
    print(f"有效性分布：{dict(val_c)}", flush=True)
    print(f"[模型={MODEL} slug={SLUG} mode=batch bsize={BSIZE}]", flush=True)


if __name__ == "__main__":
    main()
