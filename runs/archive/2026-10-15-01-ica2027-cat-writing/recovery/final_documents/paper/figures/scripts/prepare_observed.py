#!/usr/bin/env python3
"""离线整理图1/2采用的真实帖文周级计数，保留原始分类与排除账。

运行：AgentSociety2/.venv/bin/python paper/figures/scripts/prepare_observed.py
输入以第三章定稿指定的冻结矩阵为准；原始已标注工作簿只用于核对来源。
"""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
OUTPUT_DIR = ROOT / "paper/figures/source_data"
MATRIX = ROOT / "hypothesis_4/experiment_1/results/platform_heterogeneity/platform_weekly_matrices.json"
# 采用时的指纹已冻结在此，不依赖活跃run的路径来重新整理正式图件。
EXPECTED_MATRIX_SHA256 = "c8752e6cf5ef6cdfa16ff8151cc24fdd511020767df8e9caad80de5dd607d97f"
ACCEPTANCE_REFERENCE = {
    "path": "runs/2026-10-15-01-ica2027-cat-writing/section_3.accepted.md",
    "sha256": "d1a79d1332c2bacaaafe5257384fed1184d8fddaa7b3884d45a4613470f0b355",
    "role": "historical_acceptance_reference_not_runtime_input",
}
WORKBOOK = ROOT / "data/baseline/抖音微博小红书-全量已打标.xlsx"
ORIGINAL_GENERATOR = ROOT / "hypothesis_4/experiment_1/platform_heterogeneity.py"
PLATFORM_MAP = {"gsdata": "douyin", "sinamidu": "xiaohongshu"}
CATEGORIES = ("education", "mourning", "meme", "other", "marketing")
LABEL_MAP = {
    "教育观点讨论": "education",
    "事件悼念讨论": "mourning",
    "梗文化讨论": "meme",
    "其他讨论": "other",
    "借势营销": "marketing",
    "爬取噪音": "noise",
}
EVENT_WEEK = "2026-W13"


def relative(path: Path) -> str:
    return str(path.relative_to(ROOT))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_time(value: object) -> datetime | None:
    """沿用原矩阵生成器的日期解析，不将UTC时间转换为本地时区。"""
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    try:
        return datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    except ValueError:
        return None


def iso_week(value: datetime) -> str:
    year, week, _ = value.isocalendar()
    return f"{year}-W{week:02d}"


def week_start(week: str) -> str:
    year, number = week.split("-W")
    return date.fromisocalendar(int(year), int(number), 1).isoformat()


def inspect_workbook(weeks: list[str], cells: dict) -> dict:
    """重建指定窗口的六类计数，核查噪音和未映射标签，不重新编码。"""
    raw_counts = Counter()
    excluded_categories = defaultdict(Counter)
    excluded_time = Counter()
    earliest, latest = {}, {}
    raw_weeks = defaultdict(Counter)
    reconstructed = defaultdict(Counter)
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook.active
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index = {str(value): position for position, value in enumerate(headers)}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source_platform = str(row[index["source"]]).strip()
        if source_platform not in PLATFORM_MAP:
            continue
        platform = PLATFORM_MAP[source_platform]
        raw_counts[platform] += 1
        timestamp = parse_time(row[index["published_at"]])
        if timestamp is None:
            excluded_time[platform] += 1
            continue
        serialized_time = timestamp.isoformat()
        earliest[platform] = min(earliest.get(platform, serialized_time), serialized_time)
        latest[platform] = max(latest.get(platform, serialized_time), serialized_time)
        week = iso_week(timestamp)
        raw_weeks[platform][week] += 1
        if week not in weeks:
            excluded_time[platform] += 1
            continue
        original_category = str(row[index["内容类别"]]).strip()
        category = LABEL_MAP.get(original_category)
        if category is None:
            excluded_categories[platform][original_category] += 1
            continue
        reconstructed[(platform, week)][category] += 1
    workbook.close()
    for platform in PLATFORM_MAP.values():
        for week in weeks:
            if reconstructed[(platform, week)] != Counter(cells[(platform, week)]):
                raise ValueError(f"矩阵与原始标注不一致：{platform}, {week}")
    return {
        "raw_rows_by_platform": dict(raw_counts),
        "raw_rows_two_platforms": sum(raw_counts.values()),
        "published_at_min_by_platform": earliest,
        "published_at_max_by_platform": latest,
        "raw_week_counts_by_platform": {
            platform: dict(sorted(counts.items())) for platform, counts in raw_weeks.items()
        },
        "missing_time_or_outside_matrix_window_by_platform": dict(excluded_time),
        "unmapped_categories_excluded_before_matrix_by_platform": {
            platform: dict(counts) for platform, counts in excluded_categories.items()
        },
        "counts_match_original_labeled_workbook": True,
        "category_rule": "沿用原矩阵：六类标签可映射；存疑等未映射标签不计入矩阵；不是按有效性字段剔除全部无效帖。",
    }


def main() -> None:
    matrix = json.loads(MATRIX.read_text(encoding="utf-8"))
    if sha256(MATRIX) != EXPECTED_MATRIX_SHA256:
        raise ValueError("输入矩阵与第三章定稿记录的SHA256不一致")

    cells = {}
    for key, counts in matrix["cells"].items():
        labeler, source_platform, week = key.split("|")
        if labeler == "DeepSeek" and source_platform in PLATFORM_MAP:
            if set(counts) - (set(CATEGORIES) | {"noise"}):
                raise ValueError(f"出现未采用类别：{key}")
            if any(not isinstance(value, int) or value < 0 for value in counts.values()):
                raise ValueError(f"出现无效计数：{key}")
            cells[(PLATFORM_MAP[source_platform], week)] = counts
    weeks = sorted({week for _, week in cells})
    for source_platform, platform in PLATFORM_MAP.items():
        if any((platform, week) not in cells for week in weeks):
            raise ValueError(f"平台缺少周级单元：{platform}")
        if sum(sum(cells[(platform, week)].values()) for week in weeks) != matrix["platform_rows"][source_platform]:
            raise ValueError(f"平台行数与矩阵头不一致：{platform}")

    source_audit = inspect_workbook(weeks, cells)
    for week in weeks:
        cells[("pooled", week)] = dict(sum(
            (Counter(cells[(platform, week)]) for platform in PLATFORM_MAP.values()), Counter()
        ))

    rows, denominator, noise, pre_event_mourning = [], {}, {}, []
    for platform in (*PLATFORM_MAP.values(), "pooled"):
        denominator[platform], noise[platform] = 0, 0
        for week in weeks:
            counts = cells[(platform, week)]
            total = sum(counts.get(category, 0) for category in CATEGORIES)
            if total <= 0:
                raise ValueError(f"五类分母为空：{platform}, {week}")
            denominator[platform] += total
            noise[platform] += counts.get("noise", 0)
            for category in CATEGORIES:
                count = counts.get(category, 0)
                rows.append({"platform": platform, "week": week, "category": category,
                             "count": count, "total_five": total, "share": count / total})
            if week < EVENT_WEEK and counts.get("mourning", 0):
                pre_event_mourning.append({"platform": platform, "week": week,
                                          "count": counts["mourning"]})

    grouped_shares = defaultdict(float)
    for row in rows:
        grouped_shares[(row["platform"], row["week"])] += row["share"]
    share_error = max(abs(total - 1) for total in grouped_shares.values())
    if share_error > 1e-12:
        raise ValueError("周内比例未合计为1")
    peaks, last_week_meme = {}, {}
    for platform in (*PLATFORM_MAP.values(), "pooled"):
        subset = [row for row in rows if row["platform"] == platform]
        mourning_rows = [row for row in subset if row["category"] == "mourning"]
        max_count = max(row["count"] for row in mourning_rows)
        max_share = max(row["share"] for row in mourning_rows)
        peaks[platform] = {
            "count_peak_weeks": [row["week"] for row in mourning_rows if row["count"] == max_count],
            "count_peak": max_count,
            "share_peak_weeks": [row["week"] for row in mourning_rows if row["share"] == max_share],
            "share_peak": max_share,
        }
        last = next(row for row in subset if row["week"] == weeks[-1] and row["category"] == "meme")
        last_week_meme[platform] = {**last, "formatted_percent": f"{last['share']:.1%}"}

    checks = {
        "matrix_matches_accepted_sha256": True,
        "matrix_matches_original_labeled_workbook": True,
        "csv_rows": len(rows),
        "weeks": weeks,
        "week_count": len(weeks),
        "max_weekly_share_sum_error": share_error,
        "pooled_method": "先合并两平台各类别计数再计算份额，不平均平台份额。",
        "last_week_meme": last_week_meme,
        "mourning_peaks": peaks,
        "pre_event_mourning_cells_retained": pre_event_mourning,
        "pre_event_note": "事前悼念标签按原矩阵保留，没有按事件时间归零；pooled行重复表示同一记录的合并统计，不是额外记录。",
    }
    sources = {
        "purpose": "正文图1/图2采用的两平台真实帖文主分类周级数据",
        "labeler": "DeepSeek",
        "platform_mapping": PLATFORM_MAP,
        "category_order": CATEGORIES,
        "event_week": EVENT_WEEK,
        "acceptance_reference": ACCEPTANCE_REFERENCE,
        "window": {"first_iso_week": weeks[0], "last_iso_week": weeks[-1],
                   "first_week_monday": week_start(weeks[0]),
                   "last_week_monday": week_start(weeks[-1]), "week_count": len(weeks),
                   "note": "保留采用矩阵的全部周；不是W01起始。不补造W01—W04的零记录。ISO周直接由源published_at解析，沿用原生成器且不转本地时区。"},
        "sources": [{"path": relative(path), "sha256": sha256(path)}
                    for path in (MATRIX, WORKBOOK, ORIGINAL_GENERATOR)],
        "source_audit": source_audit,
        "denominator": {"categories": CATEGORIES, "five_category_rows_by_platform": denominator,
                        "noise_rows_excluded_from_denominator_by_platform": noise,
                        "note": "分母包含借势营销；爬取噪音只记排除账，不纳入五类比例或CSV。存疑的既有排除另见原表核对。"},
        "outputs": [relative(OUTPUT_DIR / filename) for filename in
                    ("observed_weekly.csv", "observed_sources.json", "observed_checks.json")],
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with (OUTPUT_DIR / "observed_weekly.csv").open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=("platform", "week", "category", "count", "total_five", "share"))
        writer.writeheader()
        writer.writerows(rows)
    for filename, value in (("observed_sources.json", sources), ("observed_checks.json", checks)):
        (OUTPUT_DIR / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"rows": len(rows), "weeks": [weeks[0], weeks[-1]],
                      "last_week_meme": last_week_meme, "mourning_peaks": peaks,
                      "pre_event_mourning_cells_retained": pre_event_mourning}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
