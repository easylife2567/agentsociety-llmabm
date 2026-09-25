#!/usr/bin/env python3
"""
从 JSONL 恢复 GLM 批量轮判定结果并装配成品 xlsx（一次性恢复工具）。

背景：batch_label_v3.py 的 label_batch 返回元组、append_jsonl 按 dict 序列化，
导致批量阶段所有记录被写成数组行 [row_num, category, reason, valid, failed]，
装配阶段 load_done() 按 dict 解析全部丢弃而中止。判定数据本身可完整恢复：
本脚本同时兼容 dict 行（v2 逐条阶段）与数组行（v3 批量阶段），按 last-wins 去重后装配。

8 条方舟内容审核拒绝行（SensitiveContentDetected，重试不可恢复）标记为：
category=其他讨论, valid=无效, reason=[内容审核拒绝,需人工复核]
"""
import json, sys
from collections import Counter
import openpyxl
from openpyxl import Workbook

ROOT = "/Users/easylife/Project/AgentSociety"
SRC = f"{ROOT}/抖音微博小红书-全量已打标.xlsx"
JSONL = f"{ROOT}/.label_results_GLM5_3flash.jsonl"
OUT = f"{ROOT}/抖音微博小红书-独立重打标_GLM5_3flash.xlsx"

# 1) 解析 JSONL（两种行格式兼容，按行号聚合）
recs_by_row = {}
n_dict = n_arr = n_bad = 0
with open(JSONL) as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        try:
            o = json.loads(line)
        except Exception:
            n_bad += 1
            continue
        if isinstance(o, dict) and "row_num" in o:
            rec = o
            n_dict += 1
        elif isinstance(o, list) and len(o) == 5 and isinstance(o[0], int):
            rec = {"row_num": o[0], "category": str(o[1]), "reason": str(o[2]),
                   "valid": str(o[3]), "failed": bool(o[4])}
            n_arr += 1
        else:
            n_bad += 1
            continue
        recs_by_row.setdefault(rec["row_num"], []).append(rec)
print(f"JSONL 解析：dict行={n_dict} 数组行={n_arr} 坏行={n_bad} 去重row_num={len(recs_by_row)}")

# 取数策略：每行取最新一条“成功”记录；无任何成功记录才视为失败
labels = {}
n_recovered = 0
for rn, rs in recs_by_row.items():
    ok = [r for r in rs if not r.get("failed")]
    if ok:
        if rs[-1].get("failed"):
            n_recovered += 1  # 末记录失败但已被更早的成功记录恢复
        labels[rn] = ok[-1]
    else:
        labels[rn] = rs[-1]  # 全失败行：保留末次失败记录，进入兜底
print(f"取数：成功覆盖={len(labels) - sum(1 for r in labels.values() if r.get('failed'))}，"
      f"靠更早成功恢复={n_recovered}，全失败待兜底={sum(1 for r in labels.values() if r.get('failed'))}")

# 2) 读源表
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
ws = wb["Sheet1"]
it = ws.iter_rows(values_only=True)
header = list(next(it))
rows = {}
rn = 1
for row in it:
    rn += 1
    if row is None or all(v is None for v in row):
        continue
    rows[rn] = row
wb.close()
print(f"源表 {len(rows)} 条")

# 3) 覆盖率检查 + 审核拒绝行兜底
still_failed = [n for n in sorted(rows) if n in labels and labels[n].get("failed")]
missing = [n for n in sorted(rows) if n not in labels]
print(f"失败待兜底 {len(still_failed)} 条：{still_failed}")
print(f"完全缺失 {len(missing)} 条")
if missing:
    print("❌ 存在完全缺失的行，中止装配")
    sys.exit(2)
MOD_REASON = "[内容审核拒绝,需人工复核]"
for n in still_failed:
    labels[n] = {"row_num": n, "category": "其他讨论", "reason": MOD_REASON,
                 "valid": "无效", "failed": False}

# 4) 装配（结构=原表57列，前三列替换为新判定）
out_wb = Workbook()
out_ws = out_wb.active
out_ws.append(header)
cat_c, val_c = Counter(), Counter()
for n in sorted(rows):
    row = rows[n]
    lab = labels[n]
    cat, valid = lab["category"], lab["valid"]
    assert cat in {"借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音"}, f"行{n}类别异常:{cat}"
    assert valid in {"有效", "无效"}, f"行{n}有效性异常:{valid}"
    out_ws.append([valid, cat, lab["reason"][:200]] + list(row)[3:])
    cat_c[cat] += 1
    val_c[valid] += 1
out_wb.save(OUT)

print(f"\n✅ 装配完成：{OUT}")
print(f"类别分布：{dict(cat_c.most_common())}")
print(f"有效性分布：{dict(val_c)}")
print(f"兜底[内容审核拒绝]行数：{len(still_failed)}")
