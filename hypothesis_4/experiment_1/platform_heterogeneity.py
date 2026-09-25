#!/usr/bin/env python3
"""分平台稳健性检验（r3 W3 整改：分平台曲线是否一致）。

问题
----
现实基准把抖音、微博、小红书三平台混合成一条曲线。若玩梗延迟爆发只由某一平台
（例如样本量占比 63% 的小红书）驱动，则"公共表征转移"就不能作为跨平台结论。
本脚本按 **平台 × 标注器** 分别重建周级类别份额，检验：

1. 各平台的玩梗轨迹是否都出现 W19 回升 / W20 起爆 / W22 峰值；
2. 各平台悼念峰值是否都在 W13；
3. 平台混合是否掩盖了平台间差异。

用法
----
    python platform_heterogeneity.py
"""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent.parent
OUT_DIR = SCRIPT_DIR / "results" / "platform_heterogeneity"
REPORT = OUT_DIR / "REPORT.md"

SOURCE_XLSX = ROOT / "抖音微博小红书-全量已打标.xlsx"
LABELERS = {
    "DeepSeek": SOURCE_XLSX,
    "doubao-seed": ROOT / "抖音微博小红书-独立重打标_Seed2_1_lite.xlsx",
    "GLM": ROOT / "抖音微博小红书-独立重打标_GLM5_3flash.xlsx",
}
CAT2TYPE = {
    "借势营销": "marketing",
    "事件悼念讨论": "mourning",
    "教育观点讨论": "education",
    "梗文化讨论": "meme",
    "其他讨论": "other",
    "爬取噪音": "noise",
}
VALID_TYPES = ["meme", "mourning", "education", "marketing", "other"]
TYPE_CN = {"meme": "玩梗", "mourning": "悼念", "education": "教育",
           "marketing": "营销", "other": "其他", "noise": "噪音"}
WINDOW_START, WINDOW_END, EVENT_WEEK = "2026-W05", "2026-W22", "2026-W13"


def week_key(value) -> str | None:
    if value is None:
        return None
    dt = value if isinstance(value, datetime) else None
    if dt is None:
        text = str(value).strip().replace("Z", "+00:00")
        if not text:
            return None
        try:
            dt = datetime.fromisoformat(text)
        except ValueError:
            return None
    iso = dt.isocalendar()
    return f"{iso.year}-W{iso.week:02d}"


def read_source() -> tuple[list[str | None], list[str]]:
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True)
    ws = wb.active
    weeks, platforms = [], []
    for row in ws.iter_rows(min_row=2, min_col=13, max_col=15, values_only=True):
        weeks.append(week_key(row[0]))
        platforms.append(str(row[2]).strip() if row[2] is not None else "unknown")
    wb.close()
    return weeks, platforms


def read_categories(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    cats = []
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        cats.append(str(row[0]).strip() if row[0] is not None else "")
    wb.close()
    return cats


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    weeks, platforms = read_source()
    print(f"[源表] {len(weeks)} 行；平台分布 {dict(Counter(platforms))}", flush=True)

    # (labeler, platform, week) -> Counter
    tables: dict[tuple[str, str, str], Counter] = defaultdict(Counter)
    row_counts: Counter = Counter()
    for labeler, path in LABELERS.items():
        cats = read_categories(path)
        for wk, plat, cat in zip(weeks, platforms, cats):
            if wk is None or wk < WINDOW_START or wk > WINDOW_END:
                continue
            t = CAT2TYPE.get(cat)
            if t is None:
                continue
            tables[(labeler, plat, wk)][t] += 1
            if labeler == "DeepSeek":
                row_counts[plat] += 1
        print(f"[{labeler}] 完成", flush=True)

    def shares(labeler: str, plat: str, wk: str) -> dict[str, float]:
        counter = tables.get((labeler, plat, wk), Counter())
        total = sum(counter.get(t, 0) for t in VALID_TYPES)
        return {t: (counter.get(t, 0) / total if total else 0.0) for t in VALID_TYPES}

    weeks_all = sorted({wk for (_, _, wk) in tables})
    plats = sorted(row_counts)

    lines = ["# 分平台稳健性检验", "",
             f"样本：{sum(row_counts.values())} 行（W05–W22，剔除存疑），平台构成"
             + "、".join(f"{p} {row_counts[p]}" for p in plats) + "。", ""]

    lines += ["## 1. 各平台 × 标注器的相位", ""]
    lines.append("| 标注器 | 平台 | n | 悼念峰值周 | 玩梗基线 | 玩梗起爆周 | 玩梗峰值周 | W22 玩梗 |")
    lines.append("|---|---|---|---|---|---|---|---|")
    for labeler in LABELERS:
        for plat in plats:
            n = sum(sum(tables.get((labeler, plat, wk), Counter()).values()) for wk in weeks_all)
            if not n:
                continue
            sh = {wk: shares(labeler, plat, wk) for wk in weeks_all}
            mourning_peak = max(weeks_all, key=lambda w: sh[w]["mourning"])
            pre = [sh[w]["meme"] for w in weeks_all if w < EVENT_WEEK]
            baseline = sum(pre) / len(pre) if pre else 0.0
            onset = next((w for w in weeks_all if w >= EVENT_WEEK and sh[w]["meme"] > baseline + 0.10), None)
            meme_peak = max(weeks_all, key=lambda w: sh[w]["meme"])
            lines.append(
                f"| {labeler} | {plat} | {n} | {mourning_peak.replace('2026-', '')} | "
                f"{baseline:.1%} | {(onset or '未触发').replace('2026-', '')} | "
                f"{meme_peak.replace('2026-', '')} | {sh[weeks_all[-1]]['meme']:.1%} |"
            )

    lines += ["", "## 2. 各平台周级玩梗份额（DeepSeek 标注）", ""]
    lines.append("| 周 | " + " | ".join(plats) + " |")
    lines.append("|---|" + "---|" * len(plats))
    for wk in weeks_all:
        lines.append(f"| {wk.replace('2026-', '')} | " + " | ".join(
            f"{shares('DeepSeek', p, wk)['meme']:.1%}" for p in plats) + " |")

    lines += ["", "## 3. 各平台周级悼念份额（DeepSeek 标注）", ""]
    lines.append("| 周 | " + " | ".join(plats) + " |")
    lines.append("|---|" + "---|" * len(plats))
    for wk in weeks_all:
        lines.append(f"| {wk.replace('2026-', '')} | " + " | ".join(
            f"{shares('DeepSeek', p, wk)['mourning']:.1%}" for p in plats) + " |")

    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    cache = {
        f"{l}|{p}|{w}": dict(tables.get((l, p, w), Counter()))
        for l in LABELERS for p in plats for w in weeks_all
    }
    (OUT_DIR / "platform_weekly_matrices.json").write_text(
        json.dumps({"platform_rows": dict(row_counts), "cells": cache}, ensure_ascii=False),
        encoding="utf-8")
    print(f"\n✅ 报告：{REPORT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
