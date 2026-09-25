#!/usr/bin/env python3
"""
独立批量打标脚本 v2（GLM 轮）- 可配置并发 + JSONL 断点 + 失败自动重试
不参考 DeepSeek / doubao 已有标签，独立判定三项（类别/依据/有效性）
判定提示词与 batch_label.py（Seed 轮）完全一致，保证多轮可比。

用法: python batch_label_v2.py <model_id> <slug> [workers]
产出: 抖音微博小红书-独立重打标_<slug>.xlsx（结构=原表57列，前三列替换为新判定）
"""
import os, sys, json, time, dotenv, openpyxl
from openpyxl import Workbook
from openai import OpenAI
from concurrent.futures import ThreadPoolExecutor, as_completed

dotenv.load_dotenv("/Users/easylife/Project/AgentSociety/.env")
API_KEY = os.environ["AGENTSOCIETY_LLM_API_KEY"]
API_BASE = os.environ["AGENTSOCIETY_LLM_API_BASE"]

MODEL = sys.argv[1]
SLUG = sys.argv[2]
WORKERS = int(sys.argv[3]) if len(sys.argv) > 3 else 16
# glm-5-3-flash 为强制思考型模型：max_tokens 需容纳思考 + 最终 JSON，否则 content 为空
MAX_TOKENS = int(os.environ.get("LABEL_MAX_TOKENS", "2000"))

ROOT = "/Users/easylife/Project/AgentSociety"
SRC = f"{ROOT}/抖音微博小红书-全量已打标.xlsx"
OUT = f"{ROOT}/抖音微博小红书-独立重打标_{SLUG}.xlsx"
JSONL = f"{ROOT}/.label_results_{SLUG}.jsonl"
PROGRESS = f"{ROOT}/.label_progress_{SLUG}.json"

print(f"[配置] 模型={MODEL} | slug={SLUG} | 并发={WORKERS} | max_tokens={MAX_TOKENS}", flush=True)
print(f"[配置] 输出={OUT}", flush=True)

CLIENT = OpenAI(api_key=API_KEY, base_url=API_BASE, timeout=120)

SYSTEM_PROMPT = """你是一个中文社交媒体内容分析助手。请对每条帖子进行三方面判定。

【内容类别 - 六选一，不允许自创】
- 借势营销：以商品、服务、账号引流、带货、二维码、价格、链接、限时优惠等为目的的推广内容；或借某热点人物/事件为特定产品/服务背书引流
- 事件悼念讨论：对逝者、事故、灾害、公共事件的缅怀、哀悼、点烛、RIP、默哀类内容；或围绕此类事件的严肃讨论
- 教育观点讨论：知识科普、观点输出、教育理念讨论、人生感悟、有信息增量的分析类内容（非营销）
- 梗文化讨论：meme、谐音梗、段子、玩梗、跟风玩梗、无实义的网络流行语、搞笑吐槽类内容
- 其他讨论：以上四类都不符合，但仍是正常有意义的帖子（如日常分享、生活记录、提问求助、情感抒发）
- 爬取噪音：乱码、纯表情、@好友无意义内容、无意义短文本、广告刷屏残留、空内容、格式损坏等无实质信息的条目

【数据有效性 - 二选一】
- 有效：包含可用于后续研究分析的实质内容
- 无效：无实质内容、纯噪音、纯营销引流无研究价值、或内容过于残缺无法分析

【判定原则】
1. 先判断是否为噪音：内容残缺/乱码/纯@人/纯表情 → 爬取噪音，无效
2. 再判断是否为营销：带货/引流/推广 → 借势营销，无效
3. 再判断是否为悼念：对逝者/事故/灾害的缅怀 → 事件悼念讨论，有效
4. 再判断是否为玩梗：meme/段子/谐音梗/跟风玩梗 → 梗文化讨论，有效
5. 再判断是否为教育观点：知识输出/观点分析/教育理念/人生感悟 → 教育观点讨论，有效
6. 以上都不符合 → 其他讨论，根据是否有实质内容判断有效性

【输出格式 - 严格 JSON，不要任何其他文字】
{"category": "类别名", "reason": "一句话依据", "valid": "有效|无效"}"""

VALID_CATS = {"借势营销", "事件悼念讨论", "教育观点讨论", "梗文化讨论", "其他讨论", "爬取噪音"}


def extract_json(txt):
    """从回复中提取严格 JSON 并做值域校验。"""
    if txt.startswith("```"):
        txt = txt.split("```")[1] if "```" in txt else txt
        if txt.startswith("json"):
            txt = txt[4:]
        txt = txt.strip()
    lo, hi = txt.find("{"), txt.rfind("}")
    if lo < 0 or hi <= lo:
        raise ValueError(f"no json: {txt[:60]}")
    obj = json.loads(txt[lo:hi + 1])
    cat = obj.get("category", "其他讨论")
    reason = str(obj.get("reason", ""))[:200]
    valid = obj.get("valid", "有效")
    if cat not in VALID_CATS:
        cat = "其他讨论"
    if valid not in {"有效", "无效"}:
        valid = "有效"
    return cat, reason, valid


def label_one(row_num, title, content, media):
    user_msg = f"【平台】{media}\n【标题】{title or ''}\n【正文】{(content or '')[:2000]}"
    for attempt in range(3):
        try:
            r = CLIENT.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_msg},
                ],
                temperature=0.1, max_tokens=MAX_TOKENS,
            )
            txt = (r.choices[0].message.content or "").strip()
            if not txt:
                raise ValueError("empty content (thinking exhausted max_tokens?)")
            cat, reason, valid = extract_json(txt)
            return row_num, cat, reason, valid, False
        except Exception as e:
            if attempt == 2:
                return row_num, "其他讨论", f"[API失败:{str(e)[:80]}]", "无效", True
            time.sleep(2 ** attempt)
    return row_num, "其他讨论", "[重试失败]", "无效", True


def load_done():
    """读 JSONL 断点，同一行号以最后一条为准。"""
    done = {}
    if os.path.exists(JSONL):
        with open(JSONL) as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    o = json.loads(line)
                    done[o["row_num"]] = o
                except Exception:
                    continue
    return done


def append_jsonl(recs):
    with open(JSONL, "a") as f:
        for rec in recs:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def run_pool(jobs, done_count, total, t0):
    """并发执行 jobs=[(row_num,row)]，结果追加 JSONL，返回本轮记录数。"""
    appended = 0
    batch_buf = []
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = [ex.submit(label_one, rn, row[idx["title"]], row[idx["content"]] or row[idx["full_content"]],
                          row[idx["media_type"]]) for rn, row in jobs]
        for fut in as_completed(futs):
            row_num, cat, reason, valid, failed = fut.result()
            batch_buf.append({"row_num": row_num, "category": cat, "reason": reason,
                              "valid": valid, "failed": failed})
            appended += 1
            done_count += 1
            if len(batch_buf) >= 50:
                append_jsonl(batch_buf)
                batch_buf = []
            if done_count % 200 == 0:
                elapsed = time.time() - t0
                rate = done_count / elapsed if elapsed > 0 else 0
                remaining = (total - done_count) / rate if rate > 0 else 0
                print(f"进度 {done_count}/{total} ({done_count*100//total}%) | "
                      f"速度 {rate:.1f} 条/秒 | 已跑 {elapsed/60:.1f} 分钟 | "
                      f"预计剩余 {remaining/60:.1f} 分钟", flush=True)
                with open(PROGRESS, "w") as f:
                    json.dump({"done_count": done_count, "total": total,
                               "model": MODEL, "slug": SLUG}, f)
    if batch_buf:
        append_jsonl(batch_buf)
    return appended, done_count


def main():
    t0 = time.time()
    print("读取源数据...", flush=True)
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {name: header.index(name) for name in ["title", "content", "full_content", "media_type"]}
    rows = {}
    rn = 1
    for row in it:
        rn += 1
        if row is None or all(v is None for v in row):
            continue
        rows[rn] = row
    wb.close()
    total = len(rows)
    print(f"共 {total} 条源数据", flush=True)

    done = load_done()
    ok = {n for n, o in done.items() if not o.get("failed")}
    print(f"断点续跑：已完成 {len(ok)} 条，其中失败待重试 {len(done) - len(ok)} 条", flush=True)
    with open(PROGRESS, "w") as f:
        json.dump({"done_count": len(ok), "total": total, "model": MODEL, "slug": SLUG}, f)

    pending = [n for n in sorted(rows) if n not in ok]
    print(f"待处理：{len(pending)} 条", flush=True)

    done_count = 0
    BATCH = WORKERS * 10
    for i in range(0, len(pending), BATCH):
        jobs = [(n, rows[n]) for n in pending[i:i + BATCH]]
        _, done_count = run_pool(jobs, done_count, len(pending), t0)

    # 失败行重试（最多 2 轮）
    for rnd in range(1, 3):
        cur = load_done()
        failed = [n for n in sorted(rows) if n in cur and cur[n].get("failed")]
        if not failed:
            break
        print(f"失败重试第 {rnd} 轮：{len(failed)} 条", flush=True)
        jobs = [(n, rows[n]) for n in failed]
        run_pool(jobs, 0, len(failed), t0)

    # 汇总装配：读源表逐行对齐，前三列替换为新判定
    print("装配输出 Excel...", flush=True)
    labels = load_done()
    missing = [n for n in rows if n not in labels or labels[n].get("failed")]
    if missing:
        print(f"❌ 有 {len(missing)} 条仍无有效判定，中止装配。前5条: {missing[:5]}", flush=True)
        sys.exit(2)

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.append(header)  # 表头与原表完全一致
    from collections import Counter
    cat_c, val_c = Counter(), Counter()
    for n in sorted(rows):
        row = rows[n]
        lab = labels[n]
        out_ws.append([lab["valid"], lab["category"], lab["reason"]] + list(row)[3:])
        cat_c[lab["category"]] += 1
        val_c[lab["valid"]] += 1
    out_wb.save(OUT)

    elapsed = time.time() - t0
    print(f"\n✅ 全部完成！总耗时 {elapsed/60:.1f} 分钟，输出：{OUT}", flush=True)
    print(f"类别分布：{dict(cat_c.most_common())}", flush=True)
    print(f"有效性分布：{dict(val_c)}", flush=True)
    print(f"[模型={MODEL} slug={SLUG}]", flush=True)


if __name__ == "__main__":
    main()
